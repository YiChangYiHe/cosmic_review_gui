import os
import re
import pandas as pd
from openpyxl import load_workbook
from utils.runtime_logger import RuntimeLogger
from PySide6.QtWidgets import (
    QDialog,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QPushButton,
    QFrame,
    QMessageBox,
    QWidget,
    QComboBox,
    QRadioButton,
    QButtonGroup,
    QFileDialog,
    QGroupBox,
    QCheckBox,
    QSlider,
    QScrollArea,
    QGridLayout,
)
from PySide6.QtCore import Qt, Signal, QByteArray, QSize, QThread
from PySide6.QtGui import (
    QDragEnterEvent,
    QDropEvent,
    QIcon,
    QFontMetrics,
    QPainter,
    QImage,
    QPixmap,
)

from ui.task_card import TaskCard
from utils.document_processor import DocumentProcessor
from utils.similarity_checker import SimilarityChecker
from utils.archive_utils import ArchiveUtils  # ✅ 导入压缩包处理工具
from utils.styles import apply_dark_title_bar
from utils.themes import get_tokens, is_dark_mode
from utils.path_utils import get_resource_path, clean_project_name
from utils.runtime_logger import log_debug, log_error, log_info, log_warn


class ExcelParseWorker(QThread):
    """后台解析 Excel 的工作线程，防止阻塞 UI"""
    parse_finished = Signal(str, dict)  # 信号: (key, 解析结果字典)

    def __init__(self, key, file_path):
        super().__init__()
        self.key = key
        self.file_path = file_path

    def run(self):
        # 【日志分流】解析日志归属项目日志
        from utils.runtime_logger import RuntimeLogger
        RuntimeLogger.set_project(self.key)
        try:
            # 调用静态解析方法
            result = UploadDialog.static_parse_excel(self.file_path)
            self.parse_finished.emit(self.key, result)
        except Exception as e:
            import traceback
            traceback.print_exc()
            # 发生错误时返回默认空结构，防止程序崩溃
            self.parse_finished.emit(self.key, {
                "file_path": self.file_path,
                "sheet_names": ["Sheet1"],
                "columns_info": {"Sheet1": []},
            })


class DownOnlyComboBox(QComboBox):
    """强制向下展开并屏蔽滑轮滚动的下拉框"""

    def wheelEvent(self, event):
        """屏蔽鼠标滑轮滚动切换选项，防止误触"""
        event.ignore()

    def showPopup(self):
        """重写弹出方法：强制向下展开、锁定最大高度，并给弹层容器
        运行时上整套主题样式（容器自带底色/滚动条规则，杜绝任何
        未被全局 QSS 命中的区域露出系统默认的白色横条）"""
        super().showPopup()
        # 获取下拉列表窗口
        popup = self.view().window()
        if popup:
            from utils.themes import get_tokens, is_dark_mode

            _t = get_tokens(is_dark_mode())
            popup.setStyleSheet(
                f"""
                * {{
                    background-color: {_t['bg_card']};
                    color: {_t['text_body']};
                }}
                QAbstractItemView {{
                    border: 1px solid {_t['border']};
                    border-radius: 8px;
                    padding: 4px;
                    outline: none;
                }}
                QScrollBar:vertical {{
                    background: transparent; width: 6px; margin: 2px; border: none;
                }}
                QScrollBar::handle:vertical {{
                    background: {_t['border']}; border-radius: 3px; min-height: 24px;
                }}
                QScrollBar::handle:vertical:hover {{ background: {_t['text_sub']}; }}
                QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                    height: 0px; width: 0px;
                }}
                QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical {{
                    background: transparent;
                }}
            """
            )
            # 计算位置：在控件正下方
            pos = self.mapToGlobal(self.rect().bottomLeft())
            popup.move(pos)

            # 强制设置下拉列表的最大高度（例如，5个选项的高度）
            # 假设每个选项高度为30px，那么5个就是150px
            max_height = 150  # 可以根据需要调整
            popup.setFixedHeight(max_height)


class ElidedLabel(QLabel):
    """自动省略超长文本的标签"""

    def __init__(self, text="", parent=None):
        super().__init__(text, parent)
        self._full_text = text

    def setText(self, text):
        self._full_text = text
        self._update_elided_text()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._update_elided_text()

    def _update_elided_text(self):
        metrics = QFontMetrics(self.font())
        # 使用当前宽度减去一些余量，防止计算精度问题导致溢出
        width = max(0, self.width() - 5)
        elided = metrics.elidedText(self._full_text, Qt.ElideRight, width)
        super().setText(elided)

    def sizeHint(self):
        # 提供一个较小的 sizeHint，允许布局将其压缩
        hint = super().sizeHint()
        hint.setWidth(100)
        return hint

    def minimumSizeHint(self):
        hint = super().minimumSizeHint()
        hint.setWidth(50)
        return hint


class FileQueueItem(QFrame):
    """文件队列项"""

    remove_clicked = Signal(str)

    def __init__(
            self, display_name, cleaned_name, has_word, has_excel, has_asset=False, parent=None
    ):
        super().__init__(parent)
        self.setProperty("class", "FileQueueItem")
        self.display_name = display_name
        self.cleaned_name = cleaned_name

        layout = QHBoxLayout(self)
        layout.setContentsMargins(12, 10, 12, 10)

        # 左侧：文件信息
        left_widget = QWidget()
        left_widget.setStyleSheet("background: transparent; border: none;")
        from PySide6.QtWidgets import QSizePolicy

        left_widget.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        left_layout = QVBoxLayout(left_widget)
        left_layout.setSpacing(6)
        left_layout.setContentsMargins(0, 0, 0, 0)

        # 文件名
        name_label = ElidedLabel(display_name)
        name_label.setStyleSheet(
            "font-weight: 600; font-size: 12px; border: none; background: transparent;"
        )
        name_label.setToolTip(display_name)  # 长文件名悬浮显示
        left_layout.addWidget(name_label)

        # 文件图标行
        icon_row = QHBoxLayout()
        icon_row.setSpacing(8)

        # Word 图标
        word_icon = QLabel("W")
        word_icon.setFixedSize(28, 28)
        word_icon.setAlignment(Qt.AlignCenter)
        word_icon.setStyleSheet(
            """
            background: #2b579a; color: white; border-radius: 4px;
            font-size: 11px; font-weight: bold;
        """
            if has_word
            else """
            background: transparent; color: palette(mid); border: 2px dashed palette(mid);
            border-radius: 4px; font-size: 11px; font-weight: bold;
        """
        )
        icon_row.addWidget(word_icon)

        # Excel 图标
        excel_icon = QLabel("X")
        excel_icon.setFixedSize(28, 28)
        excel_icon.setAlignment(Qt.AlignCenter)
        excel_icon.setStyleSheet(
            """
            background: #217346; color: white; border-radius: 4px;
            font-size: 11px; font-weight: bold;
        """
            if has_excel
            else """
            background: transparent; color: palette(mid); border: 2px dashed palette(mid);
            border-radius: 4px; font-size: 11px; font-weight: bold;
        """
        )
        icon_row.addWidget(excel_icon)

        # 资产清单图标
        asset_icon = QLabel("A")
        asset_icon.setFixedSize(28, 28)
        asset_icon.setAlignment(Qt.AlignCenter)
        asset_icon.setStyleSheet(
            """
            background: #b45309; color: white; border-radius: 4px;
            font-size: 11px; font-weight: bold;
        """
            if has_asset
            else """
            background: transparent; color: palette(mid); border: 2px dashed palette(mid);
            border-radius: 4px; font-size: 11px; font-weight: bold;
        """
        )
        asset_icon.setToolTip("资产清单")
        icon_row.addWidget(asset_icon)

        # 状态标签：只要有拆分表就可以开始
        status_text = "配对成功" if has_excel else "无法开始"
        status_color = '#10b981' if has_excel else '#ef4444'
        status_label = QLabel(status_text)
        status_label.setStyleSheet(
            f"""
            font-size: 11px;
            color: {status_color};
            font-weight: 600;
        """
        )
        icon_row.addWidget(status_label)
        icon_row.addStretch()

        left_layout.addLayout(icon_row)
        layout.addWidget(left_widget, stretch=1)

        # 右侧：删除按钮
        self.remove_btn = QPushButton()
        self.remove_btn.setFixedSize(32, 32)
        self.remove_btn.setCursor(Qt.PointingHandCursor)
        self.remove_btn.setProperty("class", "DeleteBtn")
        self.remove_btn.setToolTip("移除此文件")

        # 使用自定义的 SVG 回收站图标
        trash_icon_svg = """
        <svg xmlns="http://www.w3.org/2000/svg" width="24" height="24" viewBox="0 0 24 24" fill="none" stroke="#ef4444" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">
            <polyline points="3 6 5 6 21 6"></polyline>
            <path d="M19 6v14a2 2 0 0 1-2 2H7a2 2 0 0 1-2-2V6m3 0V4a2 2 0 0 1 2-2h4a2 2 0 0 1 2 2v2"></path>
            <line x1="10" y1="11" x2="10" y2="17"></line>
            <line x1="14" y1="11" x2="14" y2="17"></line>
        </svg>
        """
        trash_qicon = QIcon(
            QPixmap.fromImage(QImage.fromData(QByteArray(trash_icon_svg.encode())))
        )
        self.remove_btn.setIcon(trash_qicon)
        self.remove_btn.setIconSize(QSize(18, 18))

        self.remove_btn.setStyleSheet(
            """
            QPushButton {
                background: #fee2e2;
                border: 1px solid #fca5a5;
                border-radius: 16px;
                padding: 0px;
                color: transparent;
            }
            QPushButton:hover {
                background: #ef4444;
                border-color: #dc2626;
            }
            QPushButton:hover QIcon {
                /* 注意：QSS 无法直接改变 SVG 颜色，所以我们通过设置按钮背景色来强化视觉 */
            }
        """
        )
        # 为深色模式特殊控制图标颜色（如果需要）
        # 我们可以通过在 hover 时切换图标来实现更好的效果，但目前通过背景色切换已经很明显了。
        # 使用闭包捕获当时的 cleaned_name
        current_name = self.cleaned_name
        self.remove_btn.clicked.connect(lambda: self.remove_clicked.emit(current_name))
        layout.addWidget(self.remove_btn, 0, Qt.AlignRight | Qt.AlignVCenter)


class UploadAreaWidget(QFrame):
    """上传区域组件"""

    files_dropped = Signal(list)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.setFixedHeight(150)
        self.setObjectName("UploadArea")

        layout = QVBoxLayout(self)
        layout.setAlignment(Qt.AlignCenter)
        layout.setSpacing(10)

        # 图标
        icon = QLabel("📂")
        icon.setStyleSheet("font-size: 48px; background: transparent;")
        icon.setAlignment(Qt.AlignCenter)
        layout.addWidget(icon)

        # 主文本
        main_text = QLabel("点击或拖拽文件到此处")
        main_text.setStyleSheet("font-weight: bold; font-size: 12px;")
        main_text.setAlignment(Qt.AlignCenter)
        layout.addWidget(main_text)

        # 提示文本
        hint_text = QLabel("支持 .docx、.xlsx、.zip及.rar自动配对")
        hint_text.setProperty("class", "task-meta")
        hint_text.setStyleSheet("font-size: 11px;")
        hint_text.setAlignment(Qt.AlignCenter)
        layout.addWidget(hint_text)

    def mousePressEvent(self, event):
        """点击打开文件选择对话框"""
        files, _ = QFileDialog.getOpenFileNames(
            self,
            "选择文件",
            "",
            # 修改这里：添加 .doc 格式支持
            "文档文件 (*.doc *.docx *.xlsx);;Word文档 (*.doc *.docx);;Excel表格 (*.xlsx)",
        )
        if files:
            self.files_dropped.emit(files)

    def dragEnterEvent(self, event: QDragEnterEvent):
        """拖拽进入"""
        if event.mimeData().hasUrls():
            event.accept()
            event.acceptProposedAction()

    def dropEvent(self, event: QDropEvent):
        """拖拽放下"""
        if event.mimeData().hasUrls():
            event.accept()
            event.acceptProposedAction()
            files = [url.toLocalFile() for url in event.mimeData().urls()]
            self.files_dropped.emit(files)


class DragOverlay(QFrame):
    """拖拽覆盖层"""

    def __init__(self, parent):
        super().__init__(parent)
        # 允许点击穿透，但不拦截拖拽
        self.setAttribute(Qt.WA_TransparentForMouseEvents)
        self.hide()

        layout = QVBoxLayout(self)
        layout.setAlignment(Qt.AlignCenter)

        # 模拟图2效果
        self.container = QFrame()
        self.container.setStyleSheet(
            "background: rgba(0, 0, 0, 120); border-radius: 20px;"
        )
        container_layout = QVBoxLayout(self.container)
        container_layout.setContentsMargins(40, 40, 40, 40)
        container_layout.setSpacing(20)

        # 图标 (模拟图2中的文档图标)
        icon_label = QLabel("📄")
        icon_label.setStyleSheet("font-size: 64px; color: white;")
        icon_label.setAlignment(Qt.AlignCenter)
        container_layout.addWidget(icon_label)

        # 主文字
        self.text_label = QLabel("文件拖动到此处即可上传")
        self.text_label.setStyleSheet(
            "font-size: 20px; font-weight: bold; color: white;"
        )
        self.text_label.setAlignment(Qt.AlignCenter)
        container_layout.addWidget(self.text_label)

        # 副文字
        hint_label = QLabel("支持 Word、Excel 及 压缩包")
        hint_label.setStyleSheet("font-size: 14px; color: #cbd5e1;")
        hint_label.setAlignment(Qt.AlignCenter)
        container_layout.addWidget(hint_label)

        layout.addWidget(self.container)

        # 整体背景模糊透明效果 (使用 semi-transparent dark)
        self.setStyleSheet("background: rgba(15, 23, 42, 160);")

    def show_overlay(self):
        self.setGeometry(0, 0, self.parent().width(), self.parent().height())
        self.raise_()
        self.show()


class UploadDialog(QDialog):
    task_submitted = Signal(dict)  # ✅ 新增信号
    """上传任务弹窗 - 新UI外观 + 完整功能"""

    def __init__(self, parent=None, prefill_info=None):
        super().__init__(parent)
        self._prefill_info = prefill_info or {}
        self._pending_combo_prefill = None
        # [NEW] 资产清单下拉框预填（待资产清单解析完成后回填）
        self._pending_asset_prefill = None
        # [FIX] 预填期间暂停“按文件自动勾选节点”，防止异步解析回调覆盖回显的历史勾选
        self._prefill_active = bool(self._prefill_info)
        self.setWindowTitle("新建审核任务")
        self.setWindowIcon(QIcon(get_resource_path("ui/logo.ico")))
        self.setMinimumSize(900, 800)
        self.setAcceptDrops(True)  # 支持全局拖拽

        # 应用原生标题栏深色模式
        from PySide6.QtWidgets import QApplication

        is_dark = "background-color: #1f2937" in (
                QApplication.instance().styleSheet() or ""
        )
        apply_dark_title_bar(self, is_dark)

        self.file_queue = {}
        self.excel_info = {}  # 存储Excel文件信息
        self.asset_info = {}  # 存储资产清单Excel信息（独立于拆分表）
        self._parse_workers = []  # ✅ 新增：持有线程引用，防止被 GC 回收

        # 初始化拖拽覆盖层
        self.drag_overlay = DragOverlay(self)

        # 主布局
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        # 创建滚动区域
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        # 移除本地背景设置，遵循全局主题
        scroll_area.setFrameShape(QFrame.NoFrame)

        scroll_content = QWidget()
        scroll_content.setObjectName("ScrollContent")
        layout = QVBoxLayout(scroll_content)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(20)

        # ========== 1. 人天输入 + 并发数 ==========
        config_section = self._create_config_section()
        layout.addWidget(config_section)

        # ========== 2. 文件上传区域 ==========
        file_section = self._create_file_upload_section()
        layout.addWidget(file_section)

        # ========== 3. 节点选择 ==========
        node_section = self._create_node_selection_section()
        layout.addWidget(node_section)

        # ========== 4. 简单模式设置 ==========
        self.simple_section = self._create_simple_section()
        layout.addWidget(self.simple_section)
        # 根据 checkbox 状态初始化，默认为 True 所以这里应该显示
        self.simple_section.setVisible(self.simple_checkbox.isChecked())

        # ========== 5. 拆分表设置 ==========
        self.hierarchy_section = self._create_hierarchy_section()
        layout.addWidget(self.hierarchy_section)

        # ========== 5.5 资产清单匹配设置 ==========
        self.asset_section = self._create_asset_section()
        layout.addWidget(self.asset_section)
        self.asset_section.setVisible(self.asset_checkbox.isChecked())

        # ========== 5.6 数据属性重复检测设置 ==========
        self.data_attr_section = self._create_data_attr_section()
        layout.addWidget(self.data_attr_section)
        self.data_attr_section.setVisible(False)  # 默认隐藏，需要时显示

        # ========== 6. 匹配选项 ==========
        options_section = self._create_options_section()
        layout.addWidget(options_section)

        scroll_area.setWidget(scroll_content)
        main_layout.addWidget(scroll_area)

        # ========== 7. 底部按钮 ==========
        footer = self._create_footer()
        main_layout.addWidget(footer)

        # [NEW] 历史项目重新初评：预填上次的材料与选项
        self._apply_prefill()

    def dragEnterEvent(self, event: QDragEnterEvent):
        """支持全局拖拽进入"""
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
            self.drag_overlay.show_overlay()

    def dragLeaveEvent(self, event):
        """拖拽离开"""
        self.drag_overlay.hide()

    def dropEvent(self, event: QDropEvent):
        """支持全局拖拽放下"""
        self.drag_overlay.hide()
        if event.mimeData().hasUrls():
            files = [url.toLocalFile() for url in event.mimeData().urls()]
            self.handle_files(files)
            event.acceptProposedAction()

    def refresh_ui_after_file_change(self):
        """✅ 新增：文件增删后，统一刷新下拉框和配置区域显隐"""
        self.update_excel_combos()
        self.update_asset_combos()
        self.auto_update_checkboxes()
        self.on_mode_checkbox_changed(0)

    def _create_config_section(self):
        """创建配置区域（人天 + 并发数）"""
        section = QGroupBox("⚙️ 基础配置")
        layout = QVBoxLayout(section)
        layout.setSpacing(15)
        layout.setContentsMargins(25, 10, 25, 20)

        # 【新增】项目名称输入框
        name_row = QHBoxLayout()
        name_row.setAlignment(Qt.AlignVCenter)
        name_label = QLabel("项目名称:")
        name_label.setFixedWidth(110)
        name_label.setStyleSheet("font-size: 13px; font-weight: 600;")

        self.project_name_input = QLineEdit()
        self.project_name_input.setPlaceholderText("默认自动提取，可手动修改（用于报告归档和展示）")
        self.project_name_input.setFixedHeight(38)
        # 样式跟随主题（原先硬编码浅色边框，深色模式下刺眼）
        _t = get_tokens(is_dark_mode())
        self.project_name_input.setStyleSheet(
            f"""
                    QLineEdit {{ padding: 8px 12px; border: 1px solid {_t['border']}; border-radius: 8px;
                                 font-size: 13px; background: {_t['input_bg']}; color: {_t['text_body']}; }}
                    QLineEdit:focus {{ border-color: {_t['accent']}; }}
                """
        )

        name_row.addWidget(name_label)
        name_row.addWidget(self.project_name_input, stretch=1)
        layout.addLayout(name_row)

        # 人天输入
        row = QHBoxLayout()
        row.setAlignment(Qt.AlignVCenter)
        label = QLabel("线上送审人天:")
        label.setFixedWidth(110)
        label.setStyleSheet("font-size: 13px; font-weight: 600;")

        self.days_input = QLineEdit()
        self.days_input.setPlaceholderText("仅开启[送审比例]时必填，例如: 125.5")
        self.days_input.setFixedHeight(38)

        row.addWidget(label)
        row.addWidget(self.days_input, stretch=1)
        layout.addLayout(row)

        return section

    def _create_file_upload_section(self):
        """创建文件上传区域"""
        section = QGroupBox("📂 文件选择 (支持拖拽)")
        # 移除硬编码样式

        layout = QVBoxLayout(section)
        layout.setSpacing(15)
        layout.setContentsMargins(25, 12, 25, 20)

        # 上传区域
        self.upload_area = UploadAreaWidget()
        self.upload_area.files_dropped.connect(self.handle_files)
        layout.addWidget(self.upload_area)

        # 文件队列
        queue_row = QHBoxLayout()
        queue_label = QLabel("📋 已上传文件:")
        queue_label.setFixedWidth(110)
        queue_label.setStyleSheet("font-size: 13px; font-weight: 600;")
        queue_row.addWidget(queue_label)
        queue_row.addStretch(1)
        layout.addLayout(queue_row)

        # 滚动区域
        files_scroll = QScrollArea()
        files_scroll.setWidgetResizable(True)
        files_scroll.setFixedHeight(180)
        files_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        files_scroll.setStyleSheet(
            """
            QScrollArea {
                border: 1px solid palette(mid);
                border-radius: 8px;
                background-color: transparent;
            }
        """
        )

        files_content = QWidget()
        files_content.setStyleSheet("background: transparent;")
        self.queue_container = QVBoxLayout(files_content)
        self.queue_container.setSpacing(8)
        self.queue_container.setContentsMargins(10, 10, 10, 10)
        files_scroll.setWidget(files_content)

        layout.addWidget(files_scroll)

        # [NEW] 目录树txt导入（可选）：超大项目已有此前导出的 Word结构树.txt 时
        # 可直接导入，第0步跳过整篇文档解析
        tree_row = QHBoxLayout()
        self.import_tree_btn = QPushButton("📂 导入目录树txt（可选）")
        self.import_tree_btn.setFixedHeight(30)
        self.import_tree_btn.setCursor(Qt.PointingHandCursor)
        self.import_tree_btn.setToolTip(
            "此前运行导出过《XX-Word结构树.txt》时可直接导入，"
            "第0步将直接基于该文件构建层级树，无需重新解析大文档"
        )
        self.import_tree_btn.clicked.connect(self.on_import_tree_txt)
        self.tree_txt_label = ElidedLabel("未导入（将正常解析Word文档结构）")
        self.tree_txt_label.setStyleSheet(
            f"font-size: 12px; color: {get_tokens(is_dark_mode())['text_sub']};"
        )
        tree_row.addWidget(self.import_tree_btn)
        tree_row.addWidget(self.tree_txt_label, 1)
        layout.addLayout(tree_row)

        return section

    def on_import_tree_txt(self):
        """选择此前导出的 Word结构树.txt，注入任务信息供第0步直接使用"""
        path, _ = QFileDialog.getOpenFileName(
            self, "选择目录树txt文件", "", "Text files (*.txt);;All files (*)"
        )
        if path:
            self.tree_txt_path = path
            self.tree_txt_label.setText(path)
            QMessageBox.information(
                self,
                "导入成功",
                "已导入目录树txt，审核时第0步将直接使用该结构树，"
                "无需重新解析Word文档。",
            )

    def _create_node_selection_section(self):
        """创建节点选择区域(取代原匹配模式选择)"""
        section = QGroupBox()

        # 创建主垂直布局
        main_layout = QVBoxLayout(section)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        # 头部区域：左侧标题 + 右侧按钮
        header_widget = QWidget()
        section_header_layout = QHBoxLayout(header_widget)
        section_header_layout.setContentsMargins(25, 12, 25, 0)

        # 标题标签
        title_label = QLabel(" 审核节点选择 (可多选)")
        title_label.setStyleSheet("font-size: 14px; font-weight: bold;")
        section_header_layout.addWidget(title_label)

        section_header_layout.addStretch()

        # 全选/全不选按钮
        self.select_all_btn = QPushButton("全选")
        self.select_all_btn.setFixedSize(70, 28)
        self.select_all_btn.setCursor(Qt.PointingHandCursor)
        self.select_all_btn.setStyleSheet("""
            QPushButton {
                background-color: #3b82f6;
                color: white;
                border: none;
                border-radius: 4px;
                font-size: 12px;
                font-weight: 600;
            }
            QPushButton:hover {
                background-color: #2563eb;
            }
            QPushButton:pressed {
                background-color: #1d4ed8;
            }
        """)
        self.select_all_btn.clicked.connect(self.select_all_nodes)
        section_header_layout.addWidget(self.select_all_btn)

        self.deselect_all_btn = QPushButton("全不选")
        self.deselect_all_btn.setFixedSize(70, 28)
        self.deselect_all_btn.setCursor(Qt.PointingHandCursor)
        self.deselect_all_btn.setStyleSheet("""
            QPushButton {
                background-color: #64748b;
                color: white;
                border: none;
                border-radius: 4px;
                font-size: 12px;
                font-weight: 600;
            }
            QPushButton:hover {
                background-color: #475569;
            }
            QPushButton:pressed {
                background-color: #334155;
            }
        """)
        self.deselect_all_btn.clicked.connect(self.deselect_all_nodes)
        section_header_layout.addWidget(self.deselect_all_btn)

        main_layout.addWidget(header_widget)

        # 主内容布局
        content_layout = QVBoxLayout()
        content_layout.setSpacing(15)
        content_layout.setContentsMargins(25, 8, 25, 20)
        main_layout.addLayout(content_layout)

        # 节点 1-4 (基础校验)
        base_grid = QHBoxLayout()
        base_grid.setAlignment(Qt.AlignVCenter)
        base_grid.setContentsMargins(0, 5, 0, 5)

        self.check_template = QCheckBox("1. 模板校验")
        self.check_template.setChecked(True)

        self.check_empty = QCheckBox("2. 空值检查")
        self.check_empty.setChecked(True)

        self.check_ratio = QCheckBox("3. 送审比例")
        self.check_ratio.setChecked(True)

        self.check_factors = QCheckBox("4. 附加值因子")
        self.check_factors.setChecked(True)

        base_grid.addWidget(self.check_template)
        base_grid.addSpacing(25)
        base_grid.addWidget(self.check_empty)
        base_grid.addSpacing(25)
        base_grid.addWidget(self.check_ratio)
        base_grid.addSpacing(25)
        base_grid.addWidget(self.check_factors)
        base_grid.addStretch(1)
        content_layout.addLayout(base_grid)

        # 间隔线或间距
        content_layout.addSpacing(10)

        # 节点 5-6 (匹配逻辑) - 带详细说明
        match_content = QVBoxLayout()
        match_content.setSpacing(15)

        # 5. 层级模式
        h_container = QWidget()
        h_container.setStyleSheet("background: transparent;")
        h_layout = QVBoxLayout(h_container)
        h_layout.setContentsMargins(0, 0, 0, 0)
        h_layout.setSpacing(4)

        self.hierarchy_checkbox = QCheckBox("5. 层级匹配（三级模块匹配）")
        self.hierarchy_checkbox.setChecked(True)
        h_desc = QLabel("适用于：Excel 三级模块列 与 Word 4.1.1.x 分级对应")
        h_desc.setProperty("class", "task-meta")
        h_desc.setStyleSheet("margin-left: 28px; font-size: 11px;")

        h_layout.addWidget(self.hierarchy_checkbox)
        h_layout.addWidget(h_desc)
        match_content.addWidget(h_container)

        # 6. 简单模式
        s_container = QWidget()
        s_container.setStyleSheet("background: transparent;")
        s_layout = QVBoxLayout(s_container)
        s_layout.setContentsMargins(0, 0, 0, 0)
        s_layout.setSpacing(4)

        self.simple_checkbox = QCheckBox("6. 功能过程（单列单行匹配）")
        self.simple_checkbox.setChecked(True)
        s_desc = QLabel("适用于：Excel 单列直接展示所有功能过程")
        s_desc.setProperty("class", "task-meta")
        s_desc.setStyleSheet("margin-left: 28px; font-size: 11px;")

        s_layout.addWidget(self.simple_checkbox)
        s_layout.addWidget(s_desc)
        match_content.addWidget(s_container)

        # 7. 数据移动类型
        dm_container = QWidget()
        dm_container.setStyleSheet("background: transparent;")
        dm_layout = QVBoxLayout(dm_container)
        dm_layout.setContentsMargins(0, 0, 0, 0)
        dm_layout.setSpacing(4)

        self.dm_checkbox = QCheckBox("7. 功能过程数据移动类型（E开头W/X结束）")
        self.dm_checkbox.setChecked(True)
        dm_desc = QLabel("核对功能过程子项是否符合以“E”开头，“W”或“X”结束的规则")
        dm_desc.setProperty("class", "task-meta")
        dm_desc.setStyleSheet("margin-left: 28px; font-size: 11px;")

        dm_layout.addWidget(self.dm_checkbox)
        dm_layout.addWidget(dm_desc)
        match_content.addWidget(dm_container)

        # 8. 资产清单匹配
        asset_container = QWidget()
        asset_container.setStyleSheet("background: transparent;")
        asset_layout = QVBoxLayout(asset_container)
        asset_layout.setContentsMargins(0, 0, 0, 0)
        asset_layout.setSpacing(4)

        self.asset_checkbox = QCheckBox("8. 资产清单匹配")
        self.asset_checkbox.setChecked(True)
        asset_desc = QLabel(
            "将功能点拆分表与资产清单进行三级层级比对（未上传文件名包含“资产清单”的 Excel 时自动跳过）"
        )
        asset_desc.setProperty("class", "task-meta")
        asset_desc.setStyleSheet("margin-left: 28px; font-size: 11px;")

        asset_layout.addWidget(self.asset_checkbox)
        asset_layout.addWidget(asset_desc)
        match_content.addWidget(asset_container)

        # 9. 数据属性重复检测
        data_attr_container = QWidget()
        data_attr_container.setStyleSheet("background: transparent;")
        data_attr_layout = QVBoxLayout(data_attr_container)
        data_attr_layout.setContentsMargins(0, 0, 0, 0)
        data_attr_layout.setSpacing(4)

        self.data_attr_checkbox = QCheckBox("9. 数据属性重复检测")
        self.data_attr_checkbox.setChecked(True)
        data_attr_desc = QLabel(
            "检测Excel中K列（数据属性）的重复项，包括单功能过程内重复和跨功能过程重复（按ERX/EW/EX模式）"
        )
        data_attr_desc.setProperty("class", "task-meta")
        data_attr_desc.setStyleSheet("margin-left: 28px; font-size: 11px;")

        data_attr_layout.addWidget(self.data_attr_checkbox)
        data_attr_layout.addWidget(data_attr_desc)
        match_content.addWidget(data_attr_container)

        content_layout.addLayout(match_content)

        # 绑定信号，联动控制下方设置区域显隐
        self.simple_checkbox.stateChanged.connect(self.on_mode_checkbox_changed)
        self.hierarchy_checkbox.stateChanged.connect(self.on_mode_checkbox_changed)
        self.asset_checkbox.stateChanged.connect(self.on_mode_checkbox_changed)
        self.data_attr_checkbox.stateChanged.connect(self.on_mode_checkbox_changed)

        return section

    def on_mode_checkbox_changed(self, state):
        """复选框状态改变，联动控制设置区域显示"""
        self.simple_section.setVisible(self.simple_checkbox.isChecked())
        # 层级匹配或资产清单匹配都需要拆分表配置
        show_hierarchy = self.hierarchy_checkbox.isChecked() or self.asset_checkbox.isChecked()
        self.hierarchy_section.setVisible(show_hierarchy)
        self.asset_section.setVisible(self.asset_checkbox.isChecked())
        # 数据属性重复检测设置区域在勾选时显示
        self.data_attr_section.setVisible(self.data_attr_checkbox.isChecked())

    def _create_simple_section(self):
        """创建简单模式设置"""
        section = QGroupBox("📋 功能过程设置")
        layout = QVBoxLayout(section)
        layout.setSpacing(12)
        layout.setContentsMargins(25, 12, 25, 20)

        # Excel 工作表
        excel_row = QHBoxLayout()
        excel_row.setAlignment(Qt.AlignVCenter)
        excel_label = QLabel("Excel 工作表:")
        excel_label.setFixedWidth(110)
        excel_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        excel_label.setStyleSheet("font-size: 12px; font-weight: 600;")

        self.simple_excel_combo = DownOnlyComboBox()
        self.simple_excel_combo.setFixedHeight(36)
        self.simple_excel_combo.setMaxVisibleItems(5)
        self.simple_excel_combo.currentIndexChanged.connect(
            self.on_simple_sheet_changed
        )

        excel_row.addWidget(excel_label)
        excel_row.addWidget(self.simple_excel_combo, stretch=1)
        layout.addLayout(excel_row)

        # 功能点列
        func_row = QHBoxLayout()
        func_label = QLabel("功能点列:")
        func_label.setFixedWidth(110)
        func_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        func_label.setStyleSheet("font-size: 12px; font-weight: 600;")

        self.func_combo = DownOnlyComboBox()
        self.func_combo.setFixedHeight(36)
        self.func_combo.setMaxVisibleItems(5)

        func_row.addWidget(func_label)
        func_row.addWidget(self.func_combo, stretch=1)
        layout.addLayout(func_row)

        # 提示
        hint_row = QHBoxLayout()
        hint = QLabel("💡 选择包含功能点的列")
        hint.setProperty("class", "task-meta")
        hint.setStyleSheet("font-size: 11px;")

        hint_row.addWidget(hint, stretch=1)
        layout.addLayout(hint_row)

        return section

    def _create_hierarchy_section(self):
        """创建拆分表设置"""
        section = QGroupBox("📊 拆分表设置")
        layout = QVBoxLayout(section)
        layout.setSpacing(12)
        layout.setContentsMargins(25, 12, 25, 20)
        # Excel 工作表
        excel_row = QHBoxLayout()
        excel_row.setAlignment(Qt.AlignVCenter)
        excel_label = QLabel("Excel 工作表:")
        excel_label.setFixedWidth(110)
        excel_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        excel_label.setStyleSheet("font-size: 12px; font-weight: 600;")
        self.excel_sheet_combo = DownOnlyComboBox()
        self.excel_sheet_combo.setFixedHeight(36)
        self.excel_sheet_combo.setMaxVisibleItems(5)
        self.excel_sheet_combo.currentIndexChanged.connect(self.on_sheet_changed)
        excel_row.addWidget(excel_label)
        excel_row.addWidget(self.excel_sheet_combo, stretch=1)
        layout.addLayout(excel_row)
        # 一级模块列
        level1_row = QHBoxLayout()
        level1_label = QLabel("一级模块列:")
        level1_label.setFixedWidth(110)
        level1_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        level1_label.setStyleSheet("font-size: 12px; font-weight: 600;")
        self.level1_combo = DownOnlyComboBox()
        self.level1_combo.setFixedHeight(36)
        self.level1_combo.setMaxVisibleItems(5)
        level1_row.addWidget(level1_label)
        level1_row.addWidget(self.level1_combo, stretch=1)
        layout.addLayout(level1_row)
        # 二级模块列
        level2_row = QHBoxLayout()
        level2_label = QLabel("二级模块列:")
        level2_label.setFixedWidth(110)
        level2_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        level2_label.setStyleSheet("font-size: 12px; font-weight: 600;")
        self.level2_combo = DownOnlyComboBox()
        self.level2_combo.setFixedHeight(36)
        self.level2_combo.setMaxVisibleItems(5)
        level2_row.addWidget(level2_label)
        level2_row.addWidget(self.level2_combo, stretch=1)
        layout.addLayout(level2_row)
        # 三级模块列
        level3_row = QHBoxLayout()
        level3_label = QLabel("三级模块列:")
        level3_label.setFixedWidth(110)
        level3_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        level3_label.setStyleSheet("font-size: 12px; font-weight: 600;")
        self.level3_combo = DownOnlyComboBox()
        self.level3_combo.setFixedHeight(36)
        self.level3_combo.setMaxVisibleItems(5)
        level3_row.addWidget(level3_label)
        level3_row.addWidget(self.level3_combo, stretch=1)
        layout.addLayout(level3_row)
        # ================= 【新增】自动编号设置 =================
        self.auto_numbering_check = QCheckBox("强制自动编号（忽略文档原始编号）")
        self.auto_numbering_check.setToolTip(
            "默认情况下，程序会尝试识别 Word 标题中的手动编号（如 '5.1'）。\n"
            "勾选此项后，将强制使用程序生成的连续编号（如 1, 1.1, 1.1.1），忽略文档原始编号。"
        )
        self.auto_numbering_check.setStyleSheet("font-size: 12px; margin-top: 5px;")
        layout.addWidget(self.auto_numbering_check)
        # ============================================================
        # 提示
        hint_row = QHBoxLayout()
        hint = QLabel("💡 请选择包含一二三级模块的列")
        hint.setProperty("class", "task-meta")
        hint.setStyleSheet("font-size: 11px;")
        hint_row.addWidget(hint, stretch=1)
        layout.addLayout(hint_row)
        return section

    def on_sheet_changed(self, index):
        """层级模式：工作表变更时更新列下拉框和列信息显示"""
        log_debug(f'\n[SHEET-CHANGED] on_sheet_changed called, index={index}')
        if index < 0:
            return

        # 获取选中的工作表名称
        sheet_name = self.excel_sheet_combo.itemData(index)
        log_debug(f"[SHEET-CHANGED] itemData({index}) = '{sheet_name}'")
        if not sheet_name:
            sheet_name = self.excel_sheet_combo.itemText(index)
            log_debug(f"[SHEET-CHANGED] Fallback to itemText: '{sheet_name}'")
            # 从显示文本中提取工作表名称
            if "、" in sheet_name:
                sheet_name = sheet_name.split("、", 1)[1]
                log_debug(f"[SHEET-CHANGED] Extracted sheet name: '{sheet_name}'")

        # 获取第一个有效的Excel文件信息
        excel_key = None
        for key, info in self.file_queue.items():
            if info["has_excel"] and key in self.excel_info:
                excel_key = key
                break

        log_debug(f'[SHEET-CHANGED] Selected excel_key: {excel_key}')
        if not excel_key:
            log_debug('[SHEET-CHANGED] No valid Excel key, returning early')
            return

        # 获取该工作表的列信息
        excel_data = self.excel_info.get(excel_key)
        log_debug(f'[SHEET-CHANGED] Excel data keys: {(list(excel_data.keys()) if excel_data else None)}')
        if not excel_data:
            log_debug('[SHEET-CHANGED] No excel data, returning early')
            return

        columns_info = excel_data.get("columns_info", {}).get(sheet_name, [])
        log_debug(f"[SHEET-CHANGED] Columns for sheet '{sheet_name}': {len(columns_info)} columns")
        if columns_info:
            log_debug(f"[SHEET-CHANGED] Column details: {[(col['letter'], col['name']) for col in columns_info[:5]]}")

        log_debug(f"\n工作表 '{sheet_name}' 变更，更新列下拉框")
        log_info(f"列信息: {[(col['letter'], col['name']) for col in columns_info]}")

        # 更新所有列下拉框
        self.update_column_combo(self.level1_combo, columns_info, "一级模块列")
        self.update_column_combo(self.level2_combo, columns_info, "二级模块列")
        self.update_column_combo(self.level3_combo, columns_info, "三级模块列")

    def _create_options_section(self):
        """创建匹配选项"""
        section = QGroupBox("⚙️ 匹配选项")
        layout = QVBoxLayout(section)
        layout.setSpacing(12)
        layout.setContentsMargins(25, 12, 25, 20)

        # 启用模糊匹配
        fuzzy_row = QHBoxLayout()
        fuzzy_row.setAlignment(Qt.AlignVCenter)
        fuzzy_label = QLabel("启用模糊匹配:")
        fuzzy_label.setFixedWidth(110)
        fuzzy_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        fuzzy_label.setStyleSheet("font-size: 12px; font-weight: 600;")

        self.fuzzy_check = QCheckBox()
        self.fuzzy_check.setChecked(True)

        fuzzy_row.addWidget(fuzzy_label)
        fuzzy_row.addWidget(self.fuzzy_check, stretch=1)
        layout.addLayout(fuzzy_row)

        # 相似度阈值
        threshold_row = QHBoxLayout()
        threshold_row.setAlignment(Qt.AlignVCenter)
        threshold_label = QLabel("相似度阈值:")
        threshold_label.setFixedWidth(110)
        threshold_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        threshold_label.setStyleSheet("font-size: 12px; font-weight: 600;")

        self.threshold_slider = QSlider(Qt.Horizontal)
        self.threshold_slider.setMinimum(50)
        self.threshold_slider.setMaximum(100)
        self.threshold_slider.setValue(80)
        self.threshold_slider.setFixedHeight(28)

        self.threshold_value = QLabel("0.80")
        self.threshold_value.setStyleSheet(
            "font-size: 12px; font-weight: bold; min-width: 50px;"
        )
        self.threshold_slider.valueChanged.connect(
            lambda v: self.threshold_value.setText(f"{v / 100:.2f}")
        )

        threshold_row.addWidget(threshold_label)
        threshold_row.addWidget(self.threshold_slider, stretch=1)
        threshold_row.addWidget(self.threshold_value)
        layout.addLayout(threshold_row)

        return section

    def _create_footer(self):
        """创建底部按钮"""
        footer = QFrame()
        # 移除硬编码样式
        footer.setStyleSheet("border-top: 1px solid palette(mid); padding: 12px 20px;")

        layout = QHBoxLayout(footer)
        layout.setContentsMargins(0, 0, 0, 0)

        layout.addStretch()

        # 清除按钮
        clear_btn = QPushButton("🗑️ 清除")
        clear_btn.setFixedSize(90, 40)
        clear_btn.setStyleSheet(
            """
            QPushButton {
                color: #ef4444;
                border: 2px solid palette(mid);
                border-radius: 8px;
                font-size: 12px;
                font-weight: bold;
            }
            QPushButton:hover {
                background: palette(alternate-base);
                border-color: #ef4444;
            }
        """
        )
        clear_btn.clicked.connect(self.clear_all)
        layout.addWidget(clear_btn)

        # 开始审核按钮
        start_btn = QPushButton("开始审核")
        start_btn.setFixedSize(135, 40)
        start_btn.setObjectName("UploadBtn")
        start_btn.clicked.connect(self.start_review)
        layout.addWidget(start_btn)

        return footer

    def handle_files(self, files):
        """处理上传的文件"""
        try:
            # [FIX] 用户手动添加新文件后，恢复"按文件自动勾选节点"的规则
            self._prefill_active = False
            log_info('=' * 50)
            log_info('开始处理文件:')
            # 新增：预处理压缩包，将其内部文件展开到待处理列表中
            expanded_files = []
            for f in files:
                if ArchiveUtils.is_archive(f):
                    log_info(f'检测到压缩包: {os.path.basename(f)}，正在解压...')
                    extracted = ArchiveUtils.extract_archive(f)
                    expanded_files.extend(extracted)
                else:
                    expanded_files.append(f)

            # 使用展开后的文件列表进行后续处理
            for file_path in expanded_files:
                filename = os.path.basename(file_path)
                base_name = os.path.splitext(filename)[0]
                ext = os.path.splitext(filename)[1].lower()
                log_info(f'\n原始文件名: {filename}')
                log_info(f'基础名称: {base_name}')
                log_info(f'扩展名: {ext}')

                # 清理文件名：去掉前后缀
                cleaned_name = clean_project_name(base_name)
                # 【日志分流】以下为本项目相关内容，写入项目日志而非应用日志
                from utils.runtime_logger import RuntimeLogger
                RuntimeLogger.set_project(cleaned_name)
                log_info(f"清理后名称: '{cleaned_name}'")

                # 强化匹配策略：如果找不到完全一致的 Key，尝试搜寻是否有"高度相似"的 Key（连续 6 个字符相同）
                target_key = cleaned_name
                if cleaned_name not in self.file_queue:
                    for existing_key in self.file_queue.keys():
                        if self._is_fuzzy_match(cleaned_name, existing_key):
                            target_key = existing_key
                            log_info(f"检测到模糊匹配: '{cleaned_name}' 与现有项目 '{existing_key}' 自动合并")
                            break

                if target_key not in self.file_queue:
                    log_info(f'新建条目: {target_key}')
                    self.file_queue[target_key] = {
                        "has_word": False,
                        "has_excel": False,
                        "has_asset": False,
                        "original_names": {"word": None, "excel": None, "asset": None},
                        "file_paths": {"word": None, "excel": None, "asset": None},
                    }
                else:
                    log_debug(f'匹配到现有条目: {target_key}')

                # 排除需求清单
                if "需求清单" in filename:
                    log_info(f'-> 跳过需求清单: {filename}')
                    continue

                # 根据关键字识别文件类型
                if (ext == ".doc" or ext == ".docx") and "说明书" in filename:
                    log_info('-> 标记为 需求说明书')
                    self.file_queue[target_key]["has_word"] = True
                    self.file_queue[target_key]["original_names"]["word"] = filename
                    self.file_queue[target_key]["file_paths"]["word"] = file_path
                elif ext == ".xlsx" and "拆分表" in filename:
                    log_info('-> 标记为 拆分表')
                    self.file_queue[target_key]["has_excel"] = True
                    self.file_queue[target_key]["original_names"]["excel"] = filename
                    self.file_queue[target_key]["file_paths"]["excel"] = file_path
                    # ✅ 【修改点1】：启动后台线程解析，不再同步阻塞 UI
                    log_debug(f"[CALL] Starting background parse for Excel with key='{target_key}'")
                    self.parse_excel_file(target_key, file_path, is_asset=False)
                elif ext == ".xlsx" and "资产" in filename:
                    log_info('-> 标记为 资产清单')
                    self.file_queue[target_key]["has_asset"] = True
                    self.file_queue[target_key]["original_names"]["asset"] = filename
                    self.file_queue[target_key]["file_paths"]["asset"] = file_path
                    # ✅ 【修改点2】：启动后台线程解析，标记为资产。
                    # 原本这里复杂的 excel_info/asset_info 切换保护逻辑已移至 on_asset_parse_finished 回调中
                    log_debug(f"[CALL] Starting background parse for Asset with key='{target_key}'")
                    self.parse_excel_file(target_key, file_path, is_asset=True)
                else:
                    log_info(f'-> 忽略不符合条件的文件: {filename}')

                # 恢复无项目上下文（后续 UI 事件走应用日志）
                from utils.runtime_logger import RuntimeLogger
                RuntimeLogger.set_project(None)

            log_info('\n当前文件队列状态:')
            for key, value in self.file_queue.items():
                log_info(f"  '{key}': Word={value['has_word']}, Excel={value['has_excel']}, Asset={value.get('has_asset', False)}")
            log_info('=' * 50)

            # ✅ 【修改点3】：立即更新文件队列显示。
            # 注意：在异步模式下，这里不需要调用 update_excel_combos()，
            # 因为此时后台线程可能还没解析完，excel_info 里还没数据。
            # 下拉框的更新将由后台线程完成后的信号回调 (on_excel_parse_finished) 自动触发。
            self.update_queue_display()

            # ================= 【新增】自动填充项目名称 =================
            # 如果项目名称输入框存在且为空，则自动从文件队列中提取第一个名称填入
            if hasattr(self, 'project_name_input') and not self.project_name_input.text().strip():
                default_name = ""
                # 遍历文件队列，取第一个 key 作为默认项目名
                for key in self.file_queue.keys():
                    default_name = key
                    break
                if default_name:
                    self.project_name_input.setText(default_name)
            # ============================================================

        except Exception as e:
            from PySide6.QtWidgets import QMessageBox
            import traceback
            error_msg = f"处理文件时发生意外错误:\n{str(e)}\n\n{traceback.format_exc()}"
            log_error(error_msg)
            QMessageBox.critical(self, "错误", error_msg)

    def parse_excel_file(self, key, file_path, is_asset=False):
        """启动后台线程解析 Excel"""
        worker = ExcelParseWorker(key, file_path)
        # 根据是拆分表还是资产清单，连接不同的回调
        if is_asset:
            worker.parse_finished.connect(self.on_asset_parse_finished)
        else:
            worker.parse_finished.connect(self.on_excel_parse_finished)

        self._parse_workers.append(worker)
        # 线程结束后自动从列表中移除
        worker.finished.connect(lambda w=worker: self._parse_workers.remove(w))
        worker.start()

    def on_excel_parse_finished(self, key, data):
        """拆分表解析完成回调"""
        self.excel_info[key] = data
        log_debug(f'[DEBUG] Excel解析完成，开始预填下拉框')
        self.update_excel_combos()
        self.update_queue_display()
        # [NEW] 重新初评预填：Excel 解析完成后恢复工作表/列选择
        if getattr(self, "_pending_combo_prefill", None):
            try:
                self._apply_combo_prefill(self._pending_combo_prefill)
            finally:
                self._pending_combo_prefill = None

    def _select_sheet_combo(self, combo, sheet_name):
        """按下拉框 data(工作表名) 选中工作表"""
        if sheet_name is None:
            return
        idx = combo.findData(sheet_name)
        if idx >= 0:
            combo.setCurrentIndex(idx)

    def _select_column_combo(self, combo, col_index):
        """按 col_info.index 选中列"""
        if col_index is None:
            return
        for i in range(combo.count()):
            d = combo.itemData(i)
            if isinstance(d, dict) and d.get("index") == col_index:
                combo.setCurrentIndex(i)
                return

    def _apply_combo_prefill(self, pf):
        """[优化] 恢复上次的工作表与列选择（静默模式，不触发联动日志）"""
        if not pf:
            return

        # 1. 阻断信号：防止选中下拉框时触发 on_sheet_changed，避免日志刷屏和界面闪烁
        combos_to_block = [
            self.simple_excel_combo, self.excel_sheet_combo,
            self.func_combo, self.level1_combo, self.level2_combo, self.level3_combo
        ]
        # 如果有资产清单下拉框，也加入阻断
        if hasattr(self, 'asset_sheet_combo'):
            combos_to_block.append(self.asset_sheet_combo)
        if hasattr(self, 'asset_level_combos'):
            combos_to_block.extend(self.asset_level_combos)

        for c in combos_to_block:
            if c: c.blockSignals(True)

        try:
            # 2. 执行预填逻辑（此时不会触发任何槽函数）
            self._select_sheet_combo(self.simple_excel_combo, pf.get("simple_sheet"))
            self._select_sheet_combo(self.excel_sheet_combo, pf.get("hierarchy_sheet"))

            self._select_column_combo(self.func_combo, pf.get("functional_column_index"))
            self._select_column_combo(self.level1_combo, pf.get("level1_column_index"))
            self._select_column_combo(self.level2_combo, pf.get("level2_column_index"))
            self._select_column_combo(self.level3_combo, pf.get("level3_column_index"))

            # 资产清单预填
            if hasattr(self, 'asset_sheet_combo') and hasattr(self, 'asset_level_combos'):
                self._select_sheet_combo(self.asset_sheet_combo, pf.get("asset_sheet"))
                for i, combo in enumerate(self.asset_level_combos):
                    key_map = ["asset_level1_index", "asset_level2_index", "asset_level3_index"]
                    self._select_column_combo(combo, pf.get(key_map[i]))

            RuntimeLogger.log(f"[UI] 历史配置下拉框已静默恢复", level="DEBUG")

        finally:
            # 3. 恢复信号：预填完毕后，恢复联动功能
            for c in combos_to_block:
                if c: c.blockSignals(False)

    def _apply_prefill(self):
        """[NEW] 历史项目重新初评：把上次的材料与选项预填进对话框"""
        pf = getattr(self, "_prefill_info", None)
        log_info('=' * 60)
        log_debug(f'[DEBUG] 预填数据: {pf is not None}')
        log_debug(f'[DEBUG] 预填数据内容: {pf}')
        log_info('=' * 60)
        if not pf:
            return
        try:
            # ================= 【新增】回显项目名称 =================
            if hasattr(self, 'project_name_input') and pf.get('filename'):
                self.project_name_input.setText(pf['filename'])
                log_info(f"[PREFILL] 恢复项目名称: {pf['filename']}")
            # =====================================================
            # 1. 材料（Word / 拆分表 / 资产清单）
            pair = (pf.get("file_pairs") or [{}])[0]
            if not isinstance(pair, dict):
                pair = {}
            word = pair.get("word")
            excel = pair.get("excel")
            # [FIX] 资产清单兼容两种记录位置：file_pairs[].asset 优先，其次顶层 asset_excel
            asset = pair.get("asset") or pf.get("asset_excel")
            have_any = any(p and os.path.exists(str(p)) for p in (word, excel, asset))
            if have_any:
                # [FIX] 直接按槽位注入，不依赖文件名关键字识别（历史材料名未必含"说明书/拆分表"）
                from utils.path_utils import clean_project_name

                base = excel or word or asset
                key = clean_project_name(os.path.splitext(os.path.basename(str(base)))[0])
                if key not in self.file_queue:
                    self.file_queue[key] = {
                        "has_word": False,
                        "has_excel": False,
                        "has_asset": False,
                        "original_names": {"word": None, "excel": None, "asset": None},
                        "file_paths": {"word": None, "excel": None, "asset": None},
                    }
                entry = self.file_queue[key]
                if excel and os.path.exists(str(excel)):
                    entry["has_excel"] = True
                    entry["original_names"]["excel"] = os.path.basename(str(excel))
                    entry["file_paths"]["excel"] = str(excel)
                    self.parse_excel_file(key, str(excel), is_asset=False)
                if word and os.path.exists(str(word)):
                    entry["has_word"] = True
                    entry["original_names"]["word"] = os.path.basename(str(word))
                    entry["file_paths"]["word"] = str(word)
                # [FIX] 资产清单同样注入文件队列（显示"A"图标）并后台解析，
                #       否则资产清单匹配设置区不会随 asset_checkbox 一起回显
                if asset and os.path.exists(str(asset)):
                    entry["has_asset"] = True
                    entry["original_names"]["asset"] = os.path.basename(str(asset))
                    entry["file_paths"]["asset"] = str(asset)
                    self.parse_excel_file(key, str(asset), is_asset=True)
                self.update_queue_display()
                self._pending_combo_prefill = {
                    "simple_sheet": pf.get("simple_sheet"),
                    "hierarchy_sheet": pf.get("hierarchy_sheet"),
                    "functional_column_index": pf.get("functional_column_index"),
                    "level1_column_index": pf.get("level1_column_index"),
                    "level2_column_index": pf.get("level2_column_index"),
                    "level3_column_index": pf.get("level3_column_index"),
                }
                # [NEW] 资产清单工作表/层级列，待资产清单后台解析完成后回填
                self._pending_asset_prefill = {
                    "asset_sheet": pf.get("asset_sheet"),
                    "asset_level1_index": pf.get("asset_level1_index"),
                    "asset_level2_index": pf.get("asset_level2_index"),
                    "asset_level3_index": pf.get("asset_level3_index"),
                }
            # 2. 基础选项
            if str(pf.get("days", "")).strip():
                self.days_input.setText(str(pf["days"]))
            checks = {
                "check_template": "run_template",
                "check_empty": "run_empty",
                "check_ratio": "run_ratio",
                "check_factors": "run_factors",
                "hierarchy_checkbox": "run_hierarchy",
                "simple_checkbox": "run_simple",
                "dm_checkbox": "run_move",
                "asset_checkbox": "run_asset",
                "data_attr_checkbox": "run_data_attribute_check",
            }
            for attr, key in checks.items():
                cb = getattr(self, attr, None)
                if cb is not None:
                    cb.setChecked(bool(pf.get(key, False)))
            if hasattr(self, "auto_numbering_check"):
                self.auto_numbering_check.setChecked(
                    bool(pf.get("auto_numbering", False))
                )
            if hasattr(self, "fuzzy_check"):
                self.fuzzy_check.setChecked(bool(pf.get("fuzzy", True)))
            if hasattr(self, "threshold_slider"):
                try:
                    self.threshold_slider.setValue(
                        int(float(pf.get("threshold", 0.8)) * 100)
                    )
                except Exception:
                    pass
            # [NEW] 回显数据属性重复检测关键词
            kw = pf.get("data_attr_keywords")
            if kw:
                try:
                    self.data_attr_keywords_input.setText(
                        ", ".join(str(k) for k in kw)
                    )
                except Exception:
                    pass
            # [NEW] 回显此前导入的目录树txt
            tree_txt = pf.get("tree_txt")
            if tree_txt and os.path.exists(str(tree_txt)):
                self.tree_txt_path = str(tree_txt)
                self.tree_txt_label.setText(str(tree_txt))
            # [FIX] 复选框恢复完毕后，同步各设置区域的显隐状态
            self.on_mode_checkbox_changed(0)
        except Exception as e:
            log_warn(f'[WARN] 预填历史任务配置失败: {e}')
    def on_asset_parse_finished(self, key, data):
        """资产清单解析完成回调"""
        # 关键修复：先保存可能已存在的拆分表数据，防止被覆盖
        saved_excel_data = self.excel_info.get(key)
        self.asset_info[key] = data
        if saved_excel_data:
            self.excel_info[key] = saved_excel_data

        self.update_asset_combos()
        self.update_queue_display()
        # [NEW] 重新初评预填：资产清单解析完成后恢复工作表/层级列选择
        if getattr(self, "_pending_asset_prefill", None):
            try:
                self._apply_asset_combo_prefill(self._pending_asset_prefill)
            finally:
                self._pending_asset_prefill = None

    def _apply_asset_combo_prefill(self, pf):
        """[NEW] 恢复上次资产清单的工作表与一/二/三级对应列选择"""
        # 先选工作表：会触发 on_asset_sheet_changed 自动重填层级列下拉框，之后再选列
        self._select_sheet_combo(self.asset_sheet_combo, pf.get("asset_sheet"))
        combos = self.asset_level_combos
        self._select_column_combo(combos[0], pf.get("asset_level1_index"))
        self._select_column_combo(combos[1], pf.get("asset_level2_index"))
        self._select_column_combo(combos[2], pf.get("asset_level3_index"))

    @staticmethod
    def _get_header_keywords_for_sheet(sheet_name):
        """表头打分关键词（按工作表名分级）"""
        sheet_lower = str(sheet_name).lower()
        if any(kw in sheet_lower for kw in ["拆分", "功能点", "功能过程"]):
            return {
                "high": ["一级模块", "二级模块", "三级模块", "一级功能", "二级功能", "三级功能"],
                "medium": ["模块", "功能", "层级", "级别", "过程"],
                "low": ["名称", "分类", "项目", "序号", "编号", "描述", "备注", "类型", "状态", "系统", "内容"]
            }
        elif any(kw in sheet_lower for kw in ["资产", "清单", "建设", "工作量"]):
            return {
                "high": ["建设目标", "一级分类", "二级分类", "功能模块", "功能点名称", "功能点描述", "资产类别",
                         "工作量"],
                "medium": ["分类", "模块", "功能", "资产", "清单", "建设"],
                "low": ["名称", "描述", "序号", "编号", "备注", "类型", "状态", "内容", "分析", "复用", "共享"]
            }
        else:
            return {
                "high": [],
                "medium": ["模块", "功能", "名称", "分类", "级别", "层级", "项目", "序号", "编号"],
                "low": ["描述", "备注", "类型", "状态", "系统", "举措", "内容", "过程", "资产", "清单", "拆分",
                        "一级", "二级", "三级"]
            }

    _NEGATIVE_KEYWORDS = [
        "非必填", "按需", "非必要", "必填", "选填", "示例", "例如",
        "请删除", "请在正式提交时删除", "修订标识", "OPEX", "CAPEX", "填写注意",
        "如项目较小", "如项目较大", "项目名称", "Sheet表的名字不允许"
    ]

    @staticmethod
    def _score_header_row(row_values, row_idx, sheet_name=""):
        """表头行打分（与旧逻辑一致）"""
        if not row_values:
            return -1000
        keywords = UploadDialog._get_header_keywords_for_sheet(sheet_name)
        score = 0
        row_text = " ".join(row_values)
        has_level1 = any(kw in row_text for kw in ["一级模块", "一级功能"])
        has_level2 = any(kw in row_text for kw in ["二级模块", "二级功能"])
        has_level3 = any(kw in row_text for kw in ["三级模块", "三级功能"])
        if has_level1 and has_level2 and has_level3:
            score += 500
        for v in row_values:
            v_str = str(v).strip()
            for kw in UploadDialog._NEGATIVE_KEYWORDS:
                if kw in v_str:
                    score -= 120
            for kw in keywords["high"]:
                if kw in v_str:
                    score += 50
            for kw in keywords["medium"]:
                if kw in v_str:
                    score += 20
            for kw in keywords["low"]:
                if kw in v_str:
                    score += 10
        if row_idx == 0:
            score += 8
        elif row_idx < 3:
            score += 35
        elif row_idx < 6:
            score += 20
        elif row_idx < 10:
            score += 5
        else:
            score -= 20
        non_empty_count = len(row_values)
        if non_empty_count >= 6:
            score += 20
        elif non_empty_count >= 4:
            score += 12
        elif non_empty_count >= 2:
            score += 5
        return score

    @staticmethod
    def _read_merged_ranges(file_path, sheet_name):
        """[PERF] 从 XLSX 包内 XML 直读合并单元格范围（毫秒级），替代整本 load_workbook"""
        import zipfile
        import re
        from xml.etree import ElementTree as ET
        NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        try:
            with zipfile.ZipFile(file_path) as z:
                names = set(z.namelist())
                rels = {}
                if "xl/_rels/workbook.xml.rels" in names:
                    rel_root = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
                    for rel in rel_root:
                        rels[rel.get("Id")] = rel.get("Target")
                wb_root = ET.fromstring(z.read("xl/workbook.xml"))
                target = None
                for sh in wb_root.iter(f"{NS_MAIN}sheet"):
                    if sh.get("name") == sheet_name:
                        t = (rels.get(sh.get(f"{NS_R}id")) or "").lstrip("/")
                        target = t if t.startswith("xl/") else "xl/" + t
                        break
                if not target or target not in names:
                    return []
                data = z.read(target)
                m = re.search(rb"<mergeCells[^>]*>(.*?)</mergeCells>", data, re.S)
                if not m:
                    return []
                return [r.decode("ascii") for r in re.findall(rb'ref="([A-Z]+\d+:[A-Z]+\d+)"', m.group(1))]
        except Exception as e:
            log_debug(f"[SMART-merged] 直读合并范围失败（忽略，按无合并处理）: {e}")
            return []

    @staticmethod
    def _probe_sheet_header(ws, file_path, sheet_name, probe_rows=10):
        """[PERF] 流式读取前 probe_rows 行做表头探测（默认10行）。
        合并表头场景从 XML 直读合并范围并回填，全程不整本加载工作簿。
        返回 columns_info 列表（结构与旧版一致）。"""
        from openpyxl.utils.cell import range_boundaries

        # 1. 流式读取前 N 行（read_only 模式下 openpyxl 只解析到 max_row 即停）
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=probe_rows, values_only=True):
            rows.append(list(row))
        if not rows:
            rows = [[]]

        # 2. 与探测区相交的合并范围（可能扩列，上限300列）
        merged_bounds = []
        for ref in UploadDialog._read_merged_ranges(file_path, sheet_name):
            try:
                mc, mr, xc, xr = range_boundaries(ref)
            except Exception:
                continue
            if mr <= probe_rows:
                merged_bounds.append((mc, mr, xc, xr))

        num_cols = max([len(r) for r in rows] + [0])
        for (mc, _mr, xc, _xr) in merged_bounds:
            num_cols = max(num_cols, min(xc, 300))
        num_cols = min(max(num_cols, 1), 300)
        for r in rows:
            if len(r) < num_cols:
                r.extend([None] * (num_cols - len(r)))

        def _row_values(idx):
            return [str(v).strip() for v in rows[idx] if v is not None and str(v).strip() != ""]

        # 3. 原始网格打分定位表头行
        best_row_idx, best_score = 0, -1000
        for i in range(len(rows)):
            score = UploadDialog._score_header_row(_row_values(i), i, sheet_name)
            if score > best_score:
                best_score, best_row_idx = score, i

        non_empty_hdr = len(_row_values(best_row_idx))
        header_density = non_empty_hdr / max(num_cols, 1)
        need_merged_fix = header_density < 0.4 and non_empty_hdr < 5

        # 4. 合并表头：回填合并区域值后重新打分（与旧逻辑一致）
        if need_merged_fix and merged_bounds:
            for (mc, mr, xc, xr) in merged_bounds:
                top_left = None
                if mr - 1 < len(rows) and mc - 1 < len(rows[mr - 1]):
                    top_left = rows[mr - 1][mc - 1]
                for rr in range(mr, min(xr, probe_rows) + 1):
                    for cc in range(mc, min(xc, num_cols) + 1):
                        if rr - 1 < len(rows) and cc - 1 < len(rows[rr - 1]):
                            rows[rr - 1][cc - 1] = top_left
            best_row_idx, best_score = 0, -1000
            for i in range(len(rows)):
                score = UploadDialog._score_header_row(_row_values(i), i, sheet_name)
                if score > best_score:
                    best_score, best_row_idx = score, i
            log_debug(f"[SMART-merged] 工作表 '{sheet_name}' 智能定位表头在第 {best_row_idx + 1} 行 (openpyxl流式, 得分: {best_score})")
        else:
            log_debug(f"[SMART] 工作表 '{sheet_name}' 智能定位表头在第 {best_row_idx + 1} 行 (得分: {best_score})")

        # 5. 构建 columns_info（结构键与旧版完全一致）
        sample_idx = best_row_idx + 1
        column_info = []
        for i in range(num_cols):
            header_value = rows[best_row_idx][i] if best_row_idx < len(rows) else None
            is_empty = not (header_value is not None and str(header_value).strip() != "")
            if is_empty:
                for r in range(len(rows)):
                    v = rows[r][i]
                    if v is not None and str(v).strip() != "":
                        is_empty = False
                        break
            sample_data = rows[sample_idx][i] if sample_idx < len(rows) else None
            column_info.append({
                "index": i, "name": header_value, "header_row": best_row_idx,
                "source": f"第{best_row_idx + 1}行", "first_row_data": sample_data,
                "letter": UploadDialog.index_to_excel_column(i), "sample_data": sample_data,
                "is_empty": is_empty,
            })
        return column_info

    @staticmethod
    def static_parse_excel(file_path):
        """【核心解析逻辑】提取为静态方法，供子线程调用

        [PERF] v2：只流式读取前 10 行做表头探测，合并单元格范围从 XLSX 包内
        XML 直读，全程 read_only 模式 —— 不再 pandas 全量解析（旧行为：
        nrows=30 仍会读完整个 10000 行工作表），也不再整本加载工作簿
        （旧行为：read_only=False 一次构建全部工作表的单元格 DOM）。
        10000 行大文件解析从 20-30s 降至秒级。"""
        RuntimeLogger.log(f"\n>>> ENTERING static_parse_excel: file={os.path.basename(file_path)}")
        # .xls 等旧格式 openpyxl 不支持，走 legacy（pandas）路径
        if not str(file_path).lower().endswith((".xlsx", ".xlsm")):
            return UploadDialog._static_parse_excel_legacy(file_path)
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
        except Exception:
            return UploadDialog._static_parse_excel_legacy(file_path)
        try:
            sheet_names = wb.sheetnames
            log_info(f'工作表列表: {sheet_names}')
            result_data = {
                "file_path": file_path,
                "sheet_names": sheet_names,
                "columns_info": {},
            }
            for sheet_name in sheet_names:
                try:
                    result_data["columns_info"][sheet_name] = UploadDialog._probe_sheet_header(
                        wb[sheet_name], file_path, sheet_name, probe_rows=10
                    )
                except Exception as e:
                    log_error(f"  读取工作表 '{sheet_name}' 时出错: {str(e)}")
                    import traceback
                    traceback.print_exc()
            return result_data
        finally:
            try:
                wb.close()
            except Exception:
                pass

    @staticmethod
    def _static_parse_excel_legacy(file_path):
        """【legacy 兜底】pandas 全量解析路径：仅用于 .xls 旧格式或新版
        read_only 流式解析异常时的回退。逻辑与旧版 static_parse_excel 一致。"""
        try:
            RuntimeLogger.log(f"\n解析Excel文件(legacy): {file_path}", level="DEBUG")
            excel_file = pd.ExcelFile(file_path)
            sheet_names = excel_file.sheet_names
            log_info(f'工作表列表: {sheet_names}')

            result_data = {
                "file_path": file_path,
                "sheet_names": sheet_names,
                "columns_info": {},
            }

            def get_header_keywords_for_sheet(sheet_name):
                return UploadDialog._get_header_keywords_for_sheet(sheet_name)

            def get_row_score(row_values, row_idx, sheet_name=""):
                return UploadDialog._score_header_row(row_values, row_idx, sheet_name)

            for sheet_name in sheet_names:
                try:
                    # [PERF] 复用已打开的 ExcelFile（工作表解析后仅取前30行打分）
                    try:
                        df = excel_file.parse(sheet_name, nrows=30, header=None)
                    except TypeError:
                        df = pd.read_excel(file_path, sheet_name=sheet_name, nrows=30, header=None)
                    probe_max_row = len(df)
                    num_cols = len(df.columns)
                    best_row_idx = 0
                    best_score = -1000
                    for i in range(probe_max_row):
                        row_values = []
                        for j in range(num_cols):
                            v = df.iloc[i, j]
                            if v is not None and not pd.isna(v) and str(v).strip() != "":
                                row_values.append(str(v).strip())
                        score = get_row_score(row_values, i, sheet_name)
                        if score > best_score:
                            best_score = score
                            best_row_idx = i

                    header_density = len([v for v in df.iloc[best_row_idx] if
                                          v is not None and not pd.isna(v) and str(v).strip() != ""]) / max(num_cols, 1)
                    need_merged_fix = header_density < 0.4 and len([v for v in df.iloc[best_row_idx] if
                                                                    v is not None and not pd.isna(v) and str(
                                                                        v).strip() != ""]) < 5

                    if need_merged_fix:
                        wb = load_workbook(file_path, data_only=True, read_only=False)
                        ws = wb[sheet_name]
                        actual_max_col = ws.max_column or 1
                        for merged_range in ws.merged_cells.ranges:
                            if merged_range.max_col > actual_max_col: actual_max_col = merged_range.max_col
                        max_col = min(300, actual_max_col)
                        probe_max_row_wb = min(30, ws.max_row or 1)
                        cell_values = {}
                        for row in ws.iter_rows(min_row=1, max_row=probe_max_row_wb, max_col=max_col,
                                                values_only=False):
                            for cell in row:
                                cell_values[(cell.row, cell.column)] = cell.value
                        for merged_range in ws.merged_cells.ranges:
                            if merged_range.min_row > probe_max_row_wb or merged_range.min_col > max_col: continue
                            value = merged_range.start_cell.value
                            end_row = min(merged_range.max_row, probe_max_row_wb)
                            end_col = min(merged_range.max_col, max_col)
                            for r in range(merged_range.min_row, end_row + 1):
                                for c in range(merged_range.min_col, end_col + 1):
                                    cell_values[(r, c)] = value
                        wb.close()
                        best_row_idx = 0
                        best_score = -1000
                        for i in range(probe_max_row_wb):
                            row_values = []
                            for c in range(1, max_col + 1):
                                v = cell_values.get((i + 1, c))
                                if v is not None and str(v).strip() != "": row_values.append(str(v).strip())
                            score = get_row_score(row_values, i, sheet_name)
                            if score > best_score:
                                best_score = score
                                best_row_idx = i
                        detected_header_row = best_row_idx
                        num_cols = max_col
                        log_debug(f"[SMART-merged] 工作表 '{sheet_name}' 智能定位表头在第 {detected_header_row + 1} 行 (openpyxl, 得分: {best_score})")
                        column_info = []
                        for i in range(max_col):
                            col_idx = i + 1
                            header_value = cell_values.get((best_row_idx + 1, col_idx))
                            is_empty = True
                            if header_value is not None and str(header_value).strip() != "":
                                is_empty = False
                            else:
                                for r in range(1, probe_max_row_wb + 1):
                                    v = cell_values.get((r, col_idx))
                                    if v is not None and str(v).strip() != "":
                                        is_empty = False
                                        break
                            sample_data = cell_values.get((best_row_idx + 2, col_idx))
                            column_info.append({
                                "index": i, "name": header_value, "header_row": detected_header_row,
                                "source": f"第{detected_header_row + 1}行", "first_row_data": sample_data,
                                # ✅ 注意：这里改为了静态方法调用
                                "letter": UploadDialog.index_to_excel_column(i), "sample_data": sample_data,
                                "is_empty": is_empty,
                            })
                    else:
                        detected_header_row = best_row_idx
                        log_debug(f"[SMART] 工作表 '{sheet_name}' 智能定位表头在第 {detected_header_row + 1} 行 (得分: {best_score})")
                        column_info = []
                        for i in range(num_cols):
                            header_value = df.iloc[best_row_idx, i] if best_row_idx < probe_max_row else None
                            if pd.isna(header_value): header_value = None
                            is_empty = True
                            if header_value is not None and str(header_value).strip() != "":
                                is_empty = False
                            else:
                                for r in range(probe_max_row):
                                    v = df.iloc[r, i]
                                    if v is not None and not pd.isna(v) and str(v).strip() != "":
                                        is_empty = False
                                        break
                            sample_row = best_row_idx + 1
                            sample_data = df.iloc[sample_row, i] if sample_row < probe_max_row else None
                            if sample_data is not None and pd.isna(sample_data): sample_data = None
                            column_info.append({
                                "index": i, "name": header_value, "header_row": detected_header_row,
                                "source": f"第{detected_header_row + 1}行", "first_row_data": sample_data,
                                # ✅ 注意：这里改为了静态方法调用
                                "letter": UploadDialog.index_to_excel_column(i), "sample_data": sample_data,
                                "is_empty": is_empty,
                            })
                    result_data["columns_info"][sheet_name] = column_info
                except Exception as e:
                    log_error(f"  读取工作表 '{sheet_name}' 时出错: {str(e)}")
                    import traceback
                    traceback.print_exc()
            return result_data
        except Exception as e:
            log_error(f'解析Excel文件失败: {str(e)}')
            import traceback
            traceback.print_exc()
            return {
                "file_path": file_path,
                "sheet_names": ["Sheet1"],
                "columns_info": {"Sheet1": []},
            }
        finally:
            # 释放 pandas 持有的文件句柄，避免文件被占用而无法删除/移动
            try:
                excel_file.close()
            except Exception:
                pass

    @staticmethod
    def index_to_excel_column(index):
        """将列索引转换为Excel列字母 (改为静态方法)"""
        result = ""
        while index >= 0:
            remainder = index % 26
            result = chr(65 + remainder) + result
            index = index // 26 - 1
        return result

    def update_excel_combos(self):
        """更新所有Excel相关下拉框的选项"""
        RuntimeLogger.log(f"[UI] 开始更新Excel下拉框选项", level="DEBUG")
        try:
            # 1. 获取第一个有效的Excel文件信息
            excel_key = None
            for key, info in self.file_queue.items():
                if info["has_excel"] and key in self.excel_info:
                    excel_key = key
                    break

            if not excel_key:
                self.excel_sheet_combo.clear()
                self.excel_sheet_combo.addItem("无可用工作表", None)
                self.simple_excel_combo.clear()
                self.simple_excel_combo.addItem("无可用工作表", None)
                for combo in [self.level1_combo, self.level2_combo, self.level3_combo, self.func_combo]:
                    combo.clear()
                    combo.addItem("无可用列", None)
                return

            excel_data = self.excel_info.get(excel_key)
            if not excel_data:
                return

            sheet_names = excel_data.get("sheet_names", [])
            RuntimeLogger.log(f"[UI] 发现工作表: {sheet_names}", level="DEBUG")

            # 2. 填充工作表下拉框
            self.excel_sheet_combo.blockSignals(True)
            self.simple_excel_combo.blockSignals(True)

            self.excel_sheet_combo.clear()
            self.simple_excel_combo.clear()

            for i, sheet_name in enumerate(sheet_names):
                display_text = f"{i + 1}、{sheet_name}"
                self.excel_sheet_combo.addItem(display_text, sheet_name)
                self.simple_excel_combo.addItem(display_text, sheet_name)

            # 3. 决定默认选中第几个 Sheet (智能选择逻辑)
            default_sheet_idx = 0
            found_target = False
            for i, name in enumerate(sheet_names):
                if any(kw in name for kw in ["功能过程点拆分表", "功能点拆分表"]):
                    default_sheet_idx = i
                    found_target = True
                    break
            if not found_target:
                for i, name in enumerate(sheet_names):
                    if any(keyword in name for keyword in ["功能点", "拆分", "清单", "审核"]):
                        default_sheet_idx = i
                        found_target = True
                        break
            if not found_target and default_sheet_idx == 0 and len(sheet_names) >= 3:
                default_sheet_idx = 2

            # 4. 选中工作表 (此时信号被阻断，不会触发日志)
            if self.excel_sheet_combo.count() > default_sheet_idx:
                self.excel_sheet_combo.setCurrentIndex(default_sheet_idx)
            if self.simple_excel_combo.count() > default_sheet_idx:
                self.simple_excel_combo.setCurrentIndex(default_sheet_idx)

            self.excel_sheet_combo.blockSignals(False)
            self.simple_excel_combo.blockSignals(False)

            # 5. [关键修复] 手动调用 on_sheet_changed，强制生成列下拉框的选项！
            # 无论是不是预填，都必须先生成选项，否则预填也没法选
            RuntimeLogger.log(f"[UI] 手动触发列选项生成...", level="DEBUG")
            self.on_sheet_changed(default_sheet_idx)  # 生成一级/二级/三级模块列的选项
            self.on_simple_sheet_changed(default_sheet_idx)  # 生成功能点列的选项

            # 6. [关键修复] 如果是重新初评（有预填数据），最后再执行“选中”操作
            is_refill = bool(getattr(self, "_pending_combo_prefill", None))
            if is_refill:
                RuntimeLogger.log(f"[UI] 执行静默预填选中...", level="DEBUG")
                try:
                    self._apply_combo_prefill(self._pending_combo_prefill)
                finally:
                    self._pending_combo_prefill = None

        except Exception as e:
            RuntimeLogger.log(f"更新下拉框时出错: {str(e)}", level="ERROR")

    def _create_asset_section(self):
        """创建资产清单匹配设置（独立的资产清单列映射配置）"""
        section = QGroupBox("🗃️ 资产清单匹配设置")
        layout = QVBoxLayout(section)
        layout.setSpacing(12)
        layout.setContentsMargins(25, 12, 25, 20)

        # 资产清单工作表
        sheet_row = QHBoxLayout()
        sheet_row.setAlignment(Qt.AlignVCenter)
        sheet_label = QLabel("清单工作表:")
        sheet_label.setFixedWidth(110)
        sheet_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        sheet_label.setStyleSheet("font-size: 12px; font-weight: 600;")

        self.asset_sheet_combo = DownOnlyComboBox()
        self.asset_sheet_combo.setFixedHeight(36)
        self.asset_sheet_combo.setMaxVisibleItems(5)
        self.asset_sheet_combo.currentIndexChanged.connect(self.on_asset_sheet_changed)

        sheet_row.addWidget(sheet_label)
        sheet_row.addWidget(self.asset_sheet_combo, stretch=1)
        layout.addLayout(sheet_row)

        # 资产清单三个层级列
        self.asset_level_combos = []
        asset_level_defs = [
            ("一级对应列:", "资产清单中与拆分表【一级模块】对应的列（如 三级模块/一级分类）"),
            ("二级对应列:", "资产清单中与拆分表【二级模块】对应的列（如 功能模块/二级分类）"),
            ("三级对应列:", "资产清单中与拆分表【三级模块】对应的列（如 功能点名称）"),
        ]
        for label_text, tooltip in asset_level_defs:
            row = QHBoxLayout()
            label = QLabel(label_text)
            label.setFixedWidth(110)
            label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            label.setStyleSheet("font-size: 12px; font-weight: 600;")

            combo = DownOnlyComboBox()
            combo.setFixedHeight(36)
            combo.setMaxVisibleItems(5)
            combo.setToolTip(tooltip)

            row.addWidget(label)
            row.addWidget(combo, stretch=1)
            layout.addLayout(row)
            self.asset_level_combos.append(combo)

        # 提示
        hint_row = QHBoxLayout()
        hint = QLabel("💡 拆分表侧沿用【层级模式设置】中的一/二/三级模块列；资产清单侧请选择对应列")
        hint.setProperty("class", "task-meta")
        hint.setStyleSheet("font-size: 11px;")

        hint_row.addWidget(hint, stretch=1)
        layout.addLayout(hint_row)

        return section

    def _create_data_attr_section(self):
        """创建数据属性重复检测设置"""
        section = QGroupBox("🔍 数据属性重复检测设置")
        layout = QVBoxLayout(section)
        layout.setSpacing(12)
        layout.setContentsMargins(25, 12, 25, 20)

        # 关键词模式输入框
        keyword_row = QHBoxLayout()
        keyword_label = QLabel("关键词模式:")
        keyword_label.setFixedWidth(110)
        keyword_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        keyword_label.setStyleSheet("font-size: 12px; font-weight: 600;")

        self.data_attr_keywords_input = QLineEdit()
        self.data_attr_keywords_input.setFixedHeight(36)
        self.data_attr_keywords_input.setText("ERX, EW, EX")
        self.data_attr_keywords_input.setToolTip("输入关键词模式，用逗号分隔。例如：ERX, EW, EX")
        self.data_attr_keywords_input.setStyleSheet(f"""
            QLineEdit {{
                padding: 8px 12px;
                border: 2px solid {get_tokens(is_dark_mode())['border']};
                border-radius: 8px;
                font-size: 13px;
                background: {get_tokens(is_dark_mode())['input_bg']};
                color: {get_tokens(is_dark_mode())['text_body']};
            }}
            QLineEdit:focus {{
                border-color: {get_tokens(is_dark_mode())['accent']};
            }}
        """)

        keyword_row.addWidget(keyword_label)
        keyword_row.addWidget(self.data_attr_keywords_input, stretch=1)
        layout.addLayout(keyword_row)

        # 提示说明
        hint_text = QLabel(
            " 说明：系统会按这些模式检测跨功能过程的数据属性重复\n"
            "   • ERX: E(进入) → R(读取) → X(退出)\n"
            "   • EW: E(进入) → W(写入)\n"
            "   • EX: E(进入) → X(退出)"
        )
        hint_text.setProperty("class", "task-meta")
        hint_text.setStyleSheet(
            f"font-size: 11px; color: {get_tokens(is_dark_mode())['text_sub']}; line-height: 1.5;")
        hint_text.setWordWrap(True)
        layout.addWidget(hint_text)

        return section

    def _parse_data_attr_keywords(self):
        """解析用户输入的关键词模式

        Returns:
            list: 关键词列表，如 ["ERX", "EW", "EX"]
        """
        if not hasattr(self, 'data_attr_keywords_input'):
            return ["ERX", "EW", "EX"]  # 默认值

        text = self.data_attr_keywords_input.text().strip()
        if not text:
            return ["ERX", "EW", "EX"]  # 空输入使用默认值

        # 支持中英文逗号分隔
        keywords = [kw.strip().upper() for kw in text.replace('，', ',').split(',') if kw.strip()]

        if not keywords:
            return ["ERX", "EW", "EX"]  # 无效输入使用默认值

        return keywords

    def _first_asset_key(self):
        """获取第一个含资产清单文件的条目 Key"""
        for key, info in self.file_queue.items():
            if info.get("has_asset") and key in self.asset_info:
                return key
        return None

    def on_asset_sheet_changed(self, index):
        """资产清单：工作表变更时更新资产清单列下拉框"""
        log_debug(f'\n[ASSET-SHEET] on_asset_sheet_changed called, index={index}')
        if index < 0:
            return

        sheet_name = self.asset_sheet_combo.itemData(index)
        if not sheet_name:
            sheet_name = self.asset_sheet_combo.itemText(index)
            if "、" in sheet_name:
                sheet_name = sheet_name.split("、", 1)[1]
        log_debug(f"[ASSET-SHEET] Selected sheet: '{sheet_name}'")

        asset_key = self._first_asset_key()
        log_debug(f'[ASSET-SHEET] First asset key: {asset_key}')
        if not asset_key:
            log_debug('[ASSET-SHEET] No asset key found, returning early')
            return

        asset_data = self.asset_info.get(asset_key)
        log_debug(f'[ASSET-SHEET] Asset data keys: {(list(asset_data.keys()) if asset_data else None)}')
        if not asset_data:
            log_debug('[ASSET-SHEET] No asset data, returning early')
            return

        columns_info = asset_data.get("columns_info", {}).get(sheet_name, [])
        log_debug(f"[ASSET-SHEET] Columns info for sheet '{sheet_name}': {len(columns_info)} columns")
        if columns_info:
            log_debug(f"[ASSET-SHEET] Column details: {[(col['letter'], col['name'], col.get('is_empty')) for col in columns_info[:5]]}")
        log_debug(f"\n资产清单：工作表 '{sheet_name}' 变更，更新列下拉框")

        combo_names = ["一级模块列", "二级模块列", "三级模块列"]
        for combo, name in zip(self.asset_level_combos, combo_names):
            log_debug(f"[ASSET-SHEET] Updating combo '{name}', current count: {combo.count()}")
            self.update_column_combo(combo, columns_info, name)
            log_debug(f"[ASSET-SHEET] After update, combo '{name}' count: {combo.count()}")

        # 资产清单侧默认关键字 (按优先级排序，与常见资产清单列名匹配)
        keyword_groups = [
            ["一级分类", "一级模块", "建设目标"],  # 一级对应列：优先"一级分类"
            ["二级分类", "二级模块", "功能模块"],  # 二级对应列：优先"二级分类"
            ["功能模块", "三级模块", "功能点名称"],  # 三级对应列：优先"功能模块"
        ]
        for combo, keywords in zip(self.asset_level_combos, keyword_groups):
            found = False
            for kw in keywords:
                for i in range(1, combo.count()):
                    data = combo.itemData(i)
                    if data and isinstance(data, dict):
                        col_name = str(data.get("name", "")).strip()
                        if kw in col_name:
                            combo.setCurrentIndex(i)
                            found = True
                            break
                if found:
                    break

    def update_asset_combos(self):
        """资产清单文件添加后，更新资产清单工作表下拉框"""
        try:
            asset_key = self._first_asset_key()
            if not asset_key:
                # ✅ 修复：没有资产清单时，主动清空并提示
                self.asset_sheet_combo.clear()
                self.asset_sheet_combo.addItem("无可用工作表", None)
                for combo in self.asset_level_combos:
                    combo.clear()
                    combo.addItem("无可用列", None)
                return

            asset_data = self.asset_info.get(asset_key)
            if not asset_data:
                return

            sheet_names = asset_data.get("sheet_names", [])
            log_debug(f'\n更新资产清单下拉框，工作表: {sheet_names}')

            self.asset_sheet_combo.blockSignals(True)
            self.asset_sheet_combo.clear()
            for i, sheet_name in enumerate(sheet_names):
                self.asset_sheet_combo.addItem(f"{i + 1}、{sheet_name}", sheet_name)
            self.asset_sheet_combo.blockSignals(False)

            # 智能默认选择：优先包含"资产/清单"
            default_idx = 0
            for i, name in enumerate(sheet_names):
                if any(kw in name for kw in ["资产", "清单"]):
                    default_idx = i
                    break

            if self.asset_sheet_combo.count() > default_idx:
                self.asset_sheet_combo.setCurrentIndex(default_idx)
                self.on_asset_sheet_changed(default_idx)
        except Exception as e:
            log_error(f'更新资产清单下拉框时出错: {str(e)}')

    def update_column_combo(self, combo_box, columns_info, combo_name):
        """更新列下拉框选项"""
        log_debug(f"\n[COMBO-UPDATE] update_column_combo called for '{combo_name}'")
        log_debug(f'[COMBO-UPDATE] columns_info length: {len(columns_info)}')
        if columns_info:
            log_debug(f"[COMBO-UPDATE] First 3 columns: {[(col['letter'], col['name'], col.get('is_empty')) for col in columns_info[:3]]}")

        combo_box.clear()

        # 添加空选项
        combo_box.addItem(f"请选择{combo_name}", None)

        # 添加列选项
        added_count = 0
        for col_info in columns_info:
            if col_info["is_empty"]:
                continue

            letter = col_info["letter"]
            name = col_info["name"]
            sample = col_info["sample_data"]

            # 使用智能探测到的表头名称
            if name and not pd.isna(name) and str(name).strip() != "":
                display_text = f"{letter}. {name}"
            else:
                display_text = f"{letter}. 空列"

            # 🔥 限制显示文本长度，防止过长
            if len(display_text) > 25:
                display_text = display_text[:22] + "..."

            combo_box.addItem(display_text, col_info)
            added_count += 1

            # 显示智能探测到的表头
            tooltip_parts = [f"列{letter}:"]
            header_row_num = col_info.get("header_row", 0) + 1
            if not pd.isna(name) and str(name).strip() != "":
                tooltip_parts.append(f"表头(第{header_row_num}行): {name}")

            # 显示示例数据
            if sample and not pd.isna(sample):
                sample_str = str(sample)
                tooltip_parts.append(f"示例: {sample_str}")

            if len(tooltip_parts) > 1:
                combo_box.setItemData(
                    combo_box.count() - 1, "\n".join(tooltip_parts), Qt.ToolTipRole
                )
            else:
                combo_box.setItemData(
                    combo_box.count() - 1, f"列{letter}: 无数据", Qt.ToolTipRole
                )

        log_debug(f'[COMBO-UPDATE] {combo_name} 下拉框已更新，共 {added_count} 个有效选项 (total items: {combo_box.count()})')

        # 应用默认选中项
        if combo_box.count() > 1:  # Index 0 is "Please select..."

            target_idx = -1
            target_keywords = []

            if "功能点" in combo_name:
                target_idx = 6  # Default to 7th column (index 6, fallback)
                target_keywords = ["子过程描述", "功能过程", "功能描述", "功能点"]
            elif "一级" in combo_name:
                target_idx = 1  # Default to 2nd column (index 1)
                target_keywords = ["一级模块", "一级功能", "模块一", "功能一"]
            elif "二级" in combo_name:
                target_idx = 2  # Default to 3rd column (index 2)
                target_keywords = ["二级模块", "二级功能", "模块二", "功能二"]
            elif "三级" in combo_name:
                target_idx = 3  # Default to 4th column (index 3)
                target_keywords = ["三级模块", "三级功能", "模块三", "功能三"]

            log_debug(f'[DEBUG] auto-selecting for {combo_name}, keywords={target_keywords}, fallback_idx={target_idx}')

            found = False

            # Strategy 1: Match by Name (Priority)
            if target_keywords:
                for i in range(1, combo_box.count()):
                    data = combo_box.itemData(i)
                    if data and isinstance(data, dict):
                        col_name = str(data.get("name", "")).strip()
                        if any(kw in col_name for kw in target_keywords):
                            log_debug(f'[DEBUG] Name match found at item {i} ({col_name})')
                            combo_box.setCurrentIndex(i)
                            found = True
                            break

            # Strategy 2: Match by Index (Fallback)
            if not found and target_idx != -1:
                for i in range(1, combo_box.count()):
                    data = combo_box.itemData(i)
                    if data and isinstance(data, dict):
                        idx = data.get("index")
                        if idx == target_idx:
                            log_debug(f'[DEBUG] Index match found at item {i}')
                            combo_box.setCurrentIndex(i)
                            found = True
                            break

            # Strategy 3: Default to First Available (Last Resort)
            if not found:
                log_debug('[DEBUG] No match found, falling back to index 1')
                combo_box.setCurrentIndex(1)

    def on_simple_sheet_changed(self, index):
        """简单模式：工作表变更时更新列下拉框和列信息显示"""
        log_debug(f'\n[SIMPLE-SHEET] on_simple_sheet_changed called, index={index}')
        if index < 0:
            return

        # 获取选中的工作表名称
        sheet_name = self.simple_excel_combo.itemData(index)
        log_debug(f"[SIMPLE-SHEET] itemData({index}) = '{sheet_name}'")
        if not sheet_name:
            sheet_name = self.simple_excel_combo.itemText(index)
            log_debug(f"[SIMPLE-SHEET] Fallback to itemText: '{sheet_name}'")
            # 从显示文本中提取工作表名称
            if "、" in sheet_name:
                sheet_name = sheet_name.split("、", 1)[1]
                log_debug(f"[SIMPLE-SHEET] Extracted sheet name: '{sheet_name}'")

        # 获取第一个有效的Excel文件信息
        excel_key = None
        for key, info in self.file_queue.items():
            if info["has_excel"] and key in self.excel_info:
                excel_key = key
                break

        log_debug(f'[SIMPLE-SHEET] Selected excel_key: {excel_key}')
        if not excel_key:
            log_debug('[SIMPLE-SHEET] No valid Excel key, returning early')
            return

        # 获取该工作表的列信息
        excel_data = self.excel_info.get(excel_key)
        log_debug(f'[SIMPLE-SHEET] Excel data keys: {(list(excel_data.keys()) if excel_data else None)}')
        if not excel_data:
            log_debug('[SIMPLE-SHEET] No excel data, returning early')
            return

        columns_info = excel_data.get("columns_info", {}).get(sheet_name, [])
        log_debug(f"[SIMPLE-SHEET] Columns for sheet '{sheet_name}': {len(columns_info)} columns")
        if columns_info:
            log_debug(f"[SIMPLE-SHEET] Column details: {[(col['letter'], col['name']) for col in columns_info[:5]]}")

        log_debug(f"\n简单模式：工作表 '{sheet_name}' 变更，更新功能点列下拉框")
        log_info(f"列信息: {[(col['letter'], col['name']) for col in columns_info]}")

        # 更新功能点列下拉框
        self.update_column_combo(self.func_combo, columns_info, "功能点列")



    def _is_fuzzy_match(self, name1, name2):
        """
        核心匹配算法：滑动窗口检查。
        只要两个名字包含连续 6 个字符的相同片段，即视为匹配。
        """
        n1 = name1.strip()
        n2 = name2.strip()
        if not n1 or not n2:
            return False

        # 如果一方包含另一方，直接成功
        if n1 in n2 or n2 in n1:
            return True

        # 滑动窗口查找 6 位连续匹配
        min_len = 6
        if len(n1) < min_len or len(n2) < min_len:
            return n1 == n2  # 太短了就硬匹配

        for i in range(len(n1) - min_len + 1):
            window = n1[i: i + min_len]
            if window in n2:
                log_info(f"[匹配成功] 发现共同片段: '{window}'")
                return True
        return False

    def update_queue_display(self):
        """更新文件队列显示"""
        # 清空现有显示
        while self.queue_container.count():
            item = self.queue_container.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

        # 添加文件项
        valid_count = 0
        for cleaned_name, status in self.file_queue.items():
            # 为显示构建一个合适的文件名
            display_name = cleaned_name
            has_asset = status.get("has_asset", False)

            # 如果有原始文件名，使用原始文件名的组合
            name_parts = []
            if status["original_names"].get("word"):
                name_parts.append(os.path.splitext(status["original_names"]["word"])[0])
            if status["original_names"].get("excel"):
                name_parts.append(os.path.splitext(status["original_names"]["excel"])[0])
            if has_asset and status["original_names"].get("asset"):
                name_parts.append(os.path.splitext(status["original_names"]["asset"])[0])

            if name_parts:
                display_name = " | ".join(name_parts)

            item = FileQueueItem(
                display_name,
                cleaned_name,
                status["has_word"],
                status["has_excel"],
                has_asset,
            )
            item.remove_clicked.connect(self.remove_file)
            self.queue_container.addWidget(item)

            if status["has_word"] and status["has_excel"]:
                valid_count += 1

        self.queue_container.addStretch()

        # 智能更新校验节点勾选状态
        self.auto_update_checkboxes()

    def auto_update_checkboxes(self):
        """根据上传的文件类型自动更新校验节点勾选状态"""
        # [FIX] 重新初评预填期间，拆分表/资产清单异步解析完成都会触发本方法，
        #       不能按“当前有哪些文件”覆盖刚恢复的历史节点勾选状态
        #       （尤其 asset_checkbox 会被 has_any_asset=False 强制改回未勾选）
        if getattr(self, "_prefill_active", False):
            return
        if not self.file_queue:
            return

        # 检查是否有任意项目包含Word文档、拆分表、资产清单
        has_any_word = any(v["has_word"] for v in self.file_queue.values())
        has_any_excel = any(v["has_excel"] for v in self.file_queue.values())
        has_any_asset = any(v.get("has_asset", False) for v in self.file_queue.values())

        # 规则1: Word文档(需规) → 模板校验、附加因子、层级匹配、功能过程
        if has_any_word:
            self.check_template.setChecked(True)
            self.check_factors.setChecked(True)
            self.hierarchy_checkbox.setChecked(True)
            self.simple_checkbox.setChecked(True)

        # 规则2: 拆分表(Excel) → 空值检查、送审比例、层级匹配、功能过程、数据移动类型、资产清单匹配、数据属性重复检测
        if has_any_excel:
            self.check_empty.setChecked(True)
            self.check_ratio.setChecked(True)
            self.hierarchy_checkbox.setChecked(True)
            self.simple_checkbox.setChecked(True)
            self.dm_checkbox.setChecked(True)
            self.asset_checkbox.setChecked(has_any_asset)  # 有资产清单才勾选
            self.data_attr_checkbox.setChecked(True)

        # 规则3: 资产清单 → 资产清单匹配
        if has_any_asset and not has_any_excel:
            # 只有资产清单没有拆分表时，只勾选资产清单匹配
            self.asset_checkbox.setChecked(True)
            # 其他依赖拆分表的节点不勾选
            self.check_empty.setChecked(False)
            self.check_ratio.setChecked(False)
            self.dm_checkbox.setChecked(False)
            self.data_attr_checkbox.setChecked(False)

        # 如果没有Word文档，取消相关节点
        if not has_any_word:
            self.check_template.setChecked(False)
            self.check_factors.setChecked(False)

        # 触发设置区域显隐更新
        self.on_mode_checkbox_changed(0)

    def validate_checkboxes(self):
        """验证勾选的节点是否有必要的文件支持

        Returns:
            dict: {
                'valid': bool,
                'message': str  # 错误信息（如果valid=False）
            }
        """
        if not self.file_queue:
            return {"valid": False, "message": "请先上传文件！"}

        # 检查是否有任意项目包含Word文档、拆分表、资产清单
        has_any_word = any(v["has_word"] for v in self.file_queue.values())
        has_any_excel = any(v["has_excel"] for v in self.file_queue.values())
        has_any_asset = any(v.get("has_asset", False) for v in self.file_queue.values())

        missing_files = []

        # Word文档相关节点验证
        word_nodes = [
            (self.check_template, "1. 模板校验"),
            (self.check_factors, "4. 附加值因子"),
        ]
        for checkbox, node_name in word_nodes:
            if checkbox.isChecked() and not has_any_word:
                missing_files.append(f"【{node_name}】需要上传需求说明书(Word文档)")

        # 层级匹配和功能过程需要Word或Excel
        if self.hierarchy_checkbox.isChecked():
            if not has_any_word and not has_any_excel:
                missing_files.append("【5. 层级匹配】需要上传需求说明书或功能点拆分表")

        if self.simple_checkbox.isChecked():
            if not has_any_word and not has_any_excel:
                missing_files.append("【6. 功能过程】需要上传需求说明书或功能点拆分表")

        # 拆分表相关节点验证
        excel_nodes = [
            (self.check_empty, "2. 空值检查"),
            (self.check_ratio, "3. 送审比例"),
            (self.dm_checkbox, "7. 功能过程数据移动类型"),
            (self.data_attr_checkbox, "9. 数据属性重复检测"),
        ]
        for checkbox, node_name in excel_nodes:
            if checkbox.isChecked() and not has_any_excel:
                missing_files.append(f"【{node_name}】需要上传功能点拆分表(Excel)")

        # 资产清单匹配验证
        if self.asset_checkbox.isChecked() and not has_any_asset:
            missing_files.append("【8. 资产清单匹配】需要上传资产清单(Excel)")

        if missing_files:
            message = "以下节点缺少必要文件，请先上传对应文件：\n\n" + "\n".join(missing_files)
            return {"valid": False, "message": message}

        return {"valid": True, "message": ""}

    def select_all_nodes(self):
        """全选所有校验节点"""
        self.check_template.setChecked(True)
        self.check_empty.setChecked(True)
        self.check_ratio.setChecked(True)
        self.check_factors.setChecked(True)
        self.hierarchy_checkbox.setChecked(True)
        self.simple_checkbox.setChecked(True)
        self.dm_checkbox.setChecked(True)
        self.asset_checkbox.setChecked(True)
        self.data_attr_checkbox.setChecked(True)
        # 触发设置区域显隐更新
        self.on_mode_checkbox_changed(0)

    def deselect_all_nodes(self):
        """全不选所有校验节点"""
        self.check_template.setChecked(False)
        self.check_empty.setChecked(False)
        self.check_ratio.setChecked(False)
        self.check_factors.setChecked(False)
        self.hierarchy_checkbox.setChecked(False)
        self.simple_checkbox.setChecked(False)
        self.dm_checkbox.setChecked(False)
        self.asset_checkbox.setChecked(False)
        self.data_attr_checkbox.setChecked(False)
        # 触发设置区域显隐更新
        self.on_mode_checkbox_changed(0)

    def remove_file(self, cleaned_name):
        """删除文件（使用清理后的文件名）"""
        if cleaned_name in self.file_queue:
            # 删除前先检查是否是最后一个文件
            is_last_file = len(self.file_queue) == 1

            # 如果是最后一个文件，清理相关缓存和文件夹
            if is_last_file:
                self._cleanup_project_files(cleaned_name)

            # 从队列中删除
            del self.file_queue[cleaned_name]

            # 同时删除所有相关的解析数据
            if cleaned_name in self.excel_info:
                del self.excel_info[cleaned_name]
            if cleaned_name in self.asset_info:
                del self.asset_info[cleaned_name]

            self.update_queue_display()
            # ✅ 新增：触发 UI 联动刷新
            self.refresh_ui_after_file_change()

    def _cleanup_project_files(self, project_name):
        """清理项目相关的缓存文件和文件夹"""
        import os
        import shutil
        from utils.initial_review_report import project_output_dir

        try:
            # 1. 删除 tree-cache 文件（Word结构树缓存）
            cache_file = os.path.join(os.getcwd(), f"{project_name}-tree_cache.json")
            if os.path.exists(cache_file):
                os.remove(cache_file)
                log_info(f'[INFO] 已删除缓存文件: {cache_file}')

            # 2. 删除项目文件夹（如果存在）
            project_dir = project_output_dir(project_name)
            if project_dir and os.path.exists(project_dir):
                # 检查文件夹是否为空或只包含缓存文件
                try:
                    shutil.rmtree(project_dir)
                    log_info(f'[INFO] 已删除项目文件夹: {project_dir}')
                except Exception as e:
                    log_warn(f'[WARN] 删除项目文件夹失败: {e}')
            # 3. 清理日志文件（可选）——仅该项目自己的日志目录，不触碰应用日志
            from utils.runtime_logger import RuntimeLogger
            log_dir = RuntimeLogger.get_log_dir(project_name=project_name)
            if log_dir and os.path.exists(log_dir):
                # 删除与项目相关的日志文件
                for f in os.listdir(log_dir):
                    if project_name in f and f.endswith('.log'):
                        log_file = os.path.join(log_dir, f)
                        try:
                            os.remove(log_file)
                            log_info(f'[INFO] 已删除日志文件: {log_file}')
                        except Exception as e:
                            log_warn(f'[WARN] 删除日志文件失败: {e}')
        except Exception as e:
            log_error(f'[ERROR] 清理项目文件时出错: {e}')

    def clear_all(self):
        """清除所有"""
        # [FIX] 清空后恢复"按文件自动勾选节点"的规则
        self._prefill_active = False

        # 清理所有项目的缓存和文件夹
        for project_name in list(self.file_queue.keys()):
            self._cleanup_project_files(project_name)

        self.file_queue.clear()
        self.excel_info.clear()
        self.asset_info.clear()  # ✅ 修复：遗漏了清空资产信息
        self.update_queue_display()
        self.days_input.clear()
        # ✅ 修复：不再手动添加 ["0", "1", "2"] 等无意义数据，统一调用刷新方法重置
        self.refresh_ui_after_file_change()

    def add_new_task(self, task_info):
        """动态添加新上传的任务卡片到主页顶部"""
        from datetime import datetime

        # 判断类型标签（示例：可根据文件名或逻辑判断“结算”或“预算”）
        # 这里简单用“结算”作为默认，你可后续扩展
        type_label = "结算"  # 或根据 task_info['filename'] 判断

        # 设置边框颜色（与 mock 一致）
        border_color = "#2563eb"  # 蓝色，代表“进行中”

        # 构造步骤状态（初始时只有第1步完成）
        steps = [
            {"num": 1, "name": "文件解析", "status": "done"},
            {"num": 2, "name": "层级匹配", "status": "pending"},
            {"num": 3, "name": "模糊比对", "status": "pending"},
            {"num": 4, "name": "人天计算", "status": "pending"},
            {"num": 5, "name": "生成报告", "status": "pending"},
        ]

        # 初始日志
        timestamp = datetime.now().strftime("%H:%M:%S")
        default_log = f"✅ 任务于 {timestamp} 提交，开始解析文件..."
        logs = {
            1: f"✅ 文件解析成功（共 {len(task_info['file_pairs'])} 对文件）",
            2: "等待匹配...",
            3: "等待比对...",
            4: "等待人天计算...",
            5: "等待生成最终报告...",
        }

        # 构造 TaskCard 所需的完整 task_data
        task_data = {
            "filename": task_info["filename"],
            "type_label": type_label,
            "days": task_info["days"],
            "border_color": border_color,
            "steps": steps,
            "default_log": default_log,
            "logs": logs,
            # 保留原始信息，供 ReportDialog 使用
            "raw_task_info": task_info,
        }

        # 创建任务卡片
        card = TaskCard(task_data)

        # 插入到滚动区域的顶部
        # 假设你在 __init__ 中已保存 scroll_layout 为实例变量
        # 如果没有，请先做这一步（见下方“重要提示”）

        self.task_layout.insertWidget(0, card)  # 插入到最上方

    def get_combo_data(self, combo):
        """安全获取下拉框选中的数据（列信息）"""
        index = combo.currentIndex()
        if index <= 0:
            return None
        return combo.itemData(index)

    def start_review(self):
        """开始审核"""
        # 1. 验证勾选的节点是否有必要的文件支持
        validation_result = self.validate_checkboxes()
        if not validation_result["valid"]:
            QMessageBox.warning(
                self,
                "缺少必要文件",
                validation_result["message"]
            )
            return

        # 2. 只有在选择了"送审比例"节点时，才必须输入线上送审人天
        if self.check_ratio.isChecked() and not self.days_input.text().strip():
            QMessageBox.warning(
                self,
                "警告",
                "已选择【3. 送审比例】节点，必须输入【线上送审人天】才能进行计算！",
            )
            # 适配深色模式错误样式
            self.days_input.setStyleSheet(
                """
                QLineEdit {
                    padding: 8px 12px;
                    border: 2px solid #ef4444;
                    border-radius: 8px;
                    font-size: 13px;
                    background-color: #450a0a;
                    color: #fca5a5;
                }
                """
            )
            return

        # 重置样式（如果之前报错过）
        self.days_input.setStyleSheet("")

        # 3. 有效文件：有拆分表的条目（Word可选，资产清单可选）
        valid_files = [k for k, v in self.file_queue.items() if v["has_excel"]]
        if not valid_files:
            QMessageBox.warning(self, "警告", "没有有效的配对文件！")
            return

        # 4. 【核心修改】：确定项目名称 (优先使用用户输入框的值)
        custom_name = ""
        if hasattr(self, 'project_name_input'):
            custom_name = self.project_name_input.text().strip()

        if custom_name:
            display_name = custom_name
            log_info(f'[OK] 使用用户自定义项目名称: {display_name}')
        else:
            # 回退逻辑：从文件名提取
            word_name = self.file_queue[valid_files[0]]["original_names"]["word"]
            if word_name:
                display_name = clean_project_name(word_name)
            else:
                display_name = clean_project_name(valid_files[0])
            log_info(f'[INFO] 使用文件名提取的项目名称: {display_name}')

        # 5. 获取选中列的表头行信息（转换为 1-based）
        func_col_info = self.func_combo.currentData()
        level1_col_info = self.level1_combo.currentData()
        func_header_row = 1  # Default to row 1 (1-based)
        if isinstance(func_col_info, dict):
            func_header_row = func_col_info.get("header_row", 0) + 1  # Convert 0-based to 1-based
        hier_header_row = 1  # Default to row 1 (1-based)
        if isinstance(level1_col_info, dict):
            hier_header_row = level1_col_info.get("header_row", 0) + 1  # Convert 0-based to 1-based

        # 6. 获取选中列的索引信息
        func_col_idx = 6  # Default
        if isinstance(func_col_info, dict):
            func_col_idx = func_col_info.get("index", 6)
        l1_col_idx = 1  # Default
        if isinstance(level1_col_info, dict):
            l1_col_idx = level1_col_info.get("index", 1)
        l2_col_info = self.level2_combo.currentData()
        l2_col_idx = 2  # Default
        if isinstance(l2_col_info, dict):
            l2_col_idx = l2_col_info.get("index", 2)
        l3_col_info = self.level3_combo.currentData()
        l3_col_idx = 3  # Default
        if isinstance(l3_col_info, dict):
            l3_col_idx = l3_col_info.get("index", 3)

        # 7. 获取选中的工作表名称 (分别针对两种模式)
        hierarchy_sheet = self.excel_sheet_combo.currentData()
        simple_sheet = self.simple_excel_combo.currentData()

        # 8. 获取资产清单配置
        asset_sheet = None
        asset_level1_idx = None
        asset_level2_idx = None
        asset_level3_idx = None
        asset_header_row = None
        if self.asset_checkbox.isChecked() and hasattr(self, 'asset_sheet_combo'):
            asset_sheet = self.asset_sheet_combo.currentData()
            log_debug(f'[DEBUG] Asset sheet selected: {asset_sheet}')
            # 获取资产清单三个层级的列索引和表头行（header_row 转换为 1-based）
            for i, combo in enumerate(self.asset_level_combos):
                col_info = combo.currentData()
                if isinstance(col_info, dict):
                    col_idx = col_info.get("index")
                    header_row_0based = col_info.get("header_row", 0)
                    header_row_1based = header_row_0based + 1  # Convert 0-based to 1-based
                    if i == 0:
                        asset_level1_idx = col_idx
                        asset_header_row = header_row_1based
                    elif i == 1:
                        asset_level2_idx = col_idx
                    elif i == 2:
                        asset_level3_idx = col_idx

        # 9. 收集资产清单文件：优先取与拆分表同项目条目的，否则取全局第一个
        asset_excel_path = None
        for key, info in self.file_queue.items():
            if info.get("has_asset") and info["file_paths"].get("asset"):
                if not asset_excel_path:
                    asset_excel_path = info["file_paths"]["asset"]

        # 10. 为每个有拆分表的条目补充 asset 路径（同条目优先，其次全局兜底）
        file_pairs = []
        for key, info in self.file_queue.items():
            if info["has_excel"]:
                pair_asset = info["file_paths"].get("asset") or asset_excel_path
                file_pairs.append(
                    {
                        "word": info["file_paths"].get("word"),
                        "excel": info["file_paths"]["excel"],
                        "asset": pair_asset,
                        "excel_sheets": info.get("worksheets", []),
                    }
                )

        # 11. 构建最终的任务信息字典
        task_info = {
            "filename": display_name,  # <--- 这里使用了统一后的 display_name
            "days": self.days_input.text(),
            "file_pairs": file_pairs,
            "excel_columns": {
                "level1": self.level1_combo.currentText(),
                "level2": self.level2_combo.currentText(),
                "level3": self.level3_combo.currentText(),
                "point": self.func_combo.currentText(),
            },
            "selected_sheet": (
                hierarchy_sheet if self.hierarchy_checkbox.isChecked() else simple_sheet
            ),  # 兼容旧逻辑
            "hierarchy_sheet": hierarchy_sheet,
            "simple_sheet": simple_sheet,
            "functional_header_row": func_header_row,
            "hierarchy_header_row": hier_header_row,
            "functional_column_index": func_col_idx,  # Pass explicit index
            "level1_column_index": l1_col_idx,  # Pass explicit index
            "level2_column_index": l2_col_idx,  # Pass explicit index
            "level3_column_index": l3_col_idx,  # Pass explicit index
            "run_template": self.check_template.isChecked(),
            "run_empty": self.check_empty.isChecked(),
            "run_ratio": self.check_ratio.isChecked(),
            "run_factors": self.check_factors.isChecked(),
            "run_hierarchy": self.hierarchy_checkbox.isChecked(),
            "run_simple": self.simple_checkbox.isChecked(),
            "run_move": self.dm_checkbox.isChecked(),
            "run_asset": self.asset_checkbox.isChecked(),
            "run_data_attribute_check": self.data_attr_checkbox.isChecked(),
            "data_attr_keywords": self._parse_data_attr_keywords(),  # 从用户输入解析关键词
            "asset_excel": asset_excel_path,
            "asset_sheet": asset_sheet,
            "asset_level1_index": asset_level1_idx,
            "asset_level2_index": asset_level2_idx,
            "asset_level3_index": asset_level3_idx,
            "asset_header_row": asset_header_row,
            "mode": (
                "both"
                if (
                        self.hierarchy_checkbox.isChecked()
                        and self.simple_checkbox.isChecked()
                )
                else ("hierarchy" if self.hierarchy_checkbox.isChecked() else "simple")
            ),
            "fuzzy": self.fuzzy_check.isChecked(),
            "threshold": self.threshold_slider.value() / 100.0,
            "validation_results": [],
            # ================= 【新增】传递自动编号配置 =================
            "auto_numbering": self.auto_numbering_check.isChecked(),
            # ============================================================
            # [NEW] 导入的目录树txt（可选），供第0步跳过结构解析
            "tree_txt": getattr(self, "tree_txt_path", ""),
        }

        # 12. [优化] 先关闭对话框，立即将控制权交还主界面，避免视觉卡顿
        self.accept()

        # [NEW] 使用 QTimer.singleShot 异步发射信号，确保对话框已经从主循环中完全退出
        # 解决"点击开始审核不会马上开始"的阻塞感
        from PySide6.QtCore import QTimer
        QTimer.singleShot(50, lambda: self.task_submitted.emit(task_info))

        # 记录日志
        log_info(f'[OK] 已通过信号异步添加任务：{display_name}')


class ImageSelectionDialog(QDialog):
    """图像选择对话框 - 用于从 Word 中提取的多张图中选择一张"""

    def __init__(self, image_paths, parent=None):
        super().__init__(parent)
        self.image_paths = image_paths
        self.selected_image = None
        self.setWindowTitle("🔍 请选择正确的架构图")
        self.setFixedSize(900, 700)
        self._setup_ui()
        apply_dark_title_bar(self)

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(15)

        title = QLabel("🖼️ 发现多张图片，请选择一张作为功能架构图进行标注")
        title.setStyleSheet("font-size: 14px; font-weight: bold; color: #3b82f6;")
        layout.addWidget(title)

        # 滚动区域显示图片列表
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("background-color: transparent; border: none;")

        scroll_content = QWidget()
        scroll_content.setStyleSheet("background-color: transparent;")
        self.grid_layout = QGridLayout(scroll_content)
        self.grid_layout.setSpacing(15)

        # 创建图片卡片
        col_count = 3
        for i, img_path in enumerate(self.image_paths):
            card = QFrame()
            card.setProperty("class", "ImageSelectCard")
            card.setStyleSheet("""
                QFrame {
                    border: 2px solid #374151;
                    border-radius: 8px;
                    background-color: rgba(31, 41, 55, 0.5);
                }
                QFrame:hover {
                    border-color: #3b82f6;
                    background-color: rgba(59, 130, 246, 0.1);
                }
            """)
            card_layout = QVBoxLayout(card)

            # 图片预览
            img_label = QLabel()
            pixmap = QPixmap(img_path)
            if not pixmap.isNull():
                scaled_pixmap = pixmap.scaled(250, 180, Qt.KeepAspectRatio, Qt.SmoothTransformation)
                img_label.setPixmap(scaled_pixmap)
            img_label.setFixedSize(250, 180)
            img_label.setAlignment(Qt.AlignCenter)
            card_layout.addWidget(img_label)

            # 选择按钮
            select_btn = QPushButton("选择此图")
            select_btn.setStyleSheet("""
                QPushButton {
                    background-color: #3b82f6;
                    color: white;
                    border-radius: 4px;
                    padding: 6px;
                }
                QPushButton:hover {
                    background-color: #2563eb;
                }
            """)
            select_btn.clicked.connect(lambda checked, p=img_path: self._on_selected(p))
            card_layout.addWidget(select_btn)

            row = i // col_count
            col = i % col_count
            self.grid_layout.addWidget(card, row, col)

        scroll.setWidget(scroll_content)
        layout.addWidget(scroll)

        # 底部
        footer = QHBoxLayout()
        cancel_btn = QPushButton("取消")
        cancel_btn.setFixedSize(100, 36)
        cancel_btn.clicked.connect(self.reject)
        footer.addStretch()
        footer.addWidget(cancel_btn)
        layout.addLayout(footer)

    def _on_selected(self, path):
        self.selected_image = path
        self.accept()